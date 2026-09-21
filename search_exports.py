from __future__ import annotations

"""Isolated job store and document renderer for homepage search exports.

The Flask process only creates/reads jobs and serves completed artifacts.  A
single-concurrency worker imports this module, claims jobs, streams hits from
the already-loaded corpus, and creates one bounded document part at a time.
"""

import html
import json
import os
import re
import shutil
import sqlite3
import threading
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from hashlib import sha256
from pathlib import Path
from typing import Callable, Iterator

from runtime_env import APPDATA_DIR
from page_labels import VERSION as PAGE_LABEL_VERSION


DB_PATH = APPDATA_DIR / "search_exports.sqlite3"
ARTIFACT_ROOT = APPDATA_DIR / "search_exports"
RETENTION_HOURS = 24
SINGLE_FILE_HIT_LIMIT = 1000
PART_HIT_LIMIT = 500
REGULAR_RESULT_LIMIT = 1000
MEMBER_RESULT_LIMIT = 5000
REGULAR_HOURLY_LIMIT = 6
MEMBER_HOURLY_LIMIT = 12
MAX_ACTIVE_PER_USER = 2
MAX_ATTEMPTS = 2
MAX_RUNTIME_SECONDS = 15 * 60
MAX_OUTPUT_BYTES = 128 * 1024 * 1024
MAX_QUERY_CHARS = 500
# Bump whenever the semantic content of a result changes.  It is part of the
# reuse fingerprint, so pre-change snippet exports cannot be served for a new
# request after complete-sentence rendering is deployed.
EXPORT_CONTENT_VERSION = "complete-sentence-v2:" + PAGE_LABEL_VERSION
WORKER_HEARTBEAT_MAX_AGE_SECONDS = 30
VALID_FORMATS = {"docx", "html"}
VALID_MODES = {"exact", "cooccurrence"}
VALID_STYLES = {"gb2025", "gb2015", "zgshkx", "mkszyj"}
ACTIVE_STATUSES = {"queued", "counting", "collecting", "rendering", "packaging"}
TERMINAL_STATUSES = {"complete", "too_large", "failed", "expired"}

_INIT_LOCK = threading.Lock()
_INITIALIZED = False
_ILLEGAL_XML_CONTROL_RE = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]"
)


class SearchExportError(RuntimeError):
    pass


class RateLimitError(SearchExportError):
    pass


class QueueLimitError(SearchExportError):
    pass


class RetryableExportError(SearchExportError):
    pass


class PausedExportError(SearchExportError):
    """A healthy cooperative pause that must not consume the retry budget."""


class OutputTooLargeError(SearchExportError):
    pass


def _tighten_permissions(path: Path, mode: int) -> None:
    if os.name == "nt":
        return
    try:
        path.chmod(mode)
    except OSError:
        pass


@dataclass(frozen=True)
class ResourceSnapshot:
    available_bytes: int | None
    load_one: float | None


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def utc_text(value: datetime | None = None) -> str:
    return (value or utc_now()).replace(microsecond=0).isoformat()


def _parse_utc(value: object) -> datetime | None:
    try:
        parsed = datetime.fromisoformat(str(value or ""))
    except (TypeError, ValueError):
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _connect() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    _tighten_permissions(DB_PATH.parent, 0o700)
    conn = sqlite3.connect(DB_PATH, timeout=30.0)
    _tighten_permissions(DB_PATH, 0o600)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout=30000")
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def init_db() -> Path:
    global _INITIALIZED
    if _INITIALIZED:
        return DB_PATH
    with _INIT_LOCK:
        if _INITIALIZED:
            return DB_PATH
        with _connect() as conn:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS search_export_jobs (
                    id TEXT PRIMARY KEY,
                    user_id INTEGER NOT NULL,
                    query_text TEXT NOT NULL,
                    normalized_query TEXT NOT NULL,
                    search_mode TEXT NOT NULL,
                    scope_json TEXT NOT NULL DEFAULT 'null',
                    book_filter TEXT NOT NULL DEFAULT '',
                    citation_style TEXT NOT NULL,
                    output_format TEXT NOT NULL,
                    corpus_version TEXT NOT NULL,
                    template_version TEXT NOT NULL,
                    request_fingerprint TEXT NOT NULL,
                    base_url TEXT NOT NULL,
                    result_limit INTEGER NOT NULL,
                    status TEXT NOT NULL,
                    result_count INTEGER NOT NULL DEFAULT 0,
                    progress_done INTEGER NOT NULL DEFAULT 0,
                    progress_total INTEGER NOT NULL DEFAULT 0,
                    output_path TEXT NOT NULL DEFAULT '',
                    output_name TEXT NOT NULL DEFAULT '',
                    output_size INTEGER NOT NULL DEFAULT 0,
                    manifest_json TEXT NOT NULL DEFAULT '{}',
                    error TEXT NOT NULL DEFAULT '',
                    attempts INTEGER NOT NULL DEFAULT 0,
                    lease_owner TEXT NOT NULL DEFAULT '',
                    lease_expires_at TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    expires_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_search_export_user_created
                    ON search_export_jobs(user_id, created_at DESC);
                CREATE INDEX IF NOT EXISTS idx_search_export_queue
                    ON search_export_jobs(status, created_at);
                CREATE INDEX IF NOT EXISTS idx_search_export_fingerprint
                    ON search_export_jobs(user_id, request_fingerprint, expires_at);

                CREATE TABLE IF NOT EXISTS search_export_settings (
                    key TEXT PRIMARY KEY,
                    value TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS search_export_worker_state (
                    worker_id TEXT PRIMARY KEY,
                    heartbeat_at TEXT NOT NULL,
                    current_job_id TEXT NOT NULL DEFAULT ''
                );
                """
            )
        ARTIFACT_ROOT.mkdir(parents=True, exist_ok=True, mode=0o700)
        _tighten_permissions(ARTIFACT_ROOT, 0o700)
        _INITIALIZED = True
    return DB_PATH


def feature_enabled() -> bool:
    init_db()
    with _connect() as conn:
        row = conn.execute(
            "SELECT value FROM search_export_settings WHERE key='enabled'"
        ).fetchone()
    if row is not None:
        return str(row["value"] or "").strip().lower() in {"1", "true", "yes", "on"}
    return str(os.environ.get("SEARCH_EXPORTS_DEFAULT_ENABLED", "0")).strip().lower() in {
        "1", "true", "yes", "on"
    }


def set_feature_enabled(enabled: bool) -> bool:
    init_db()
    now = utc_text()
    with _connect() as conn:
        conn.execute(
            """INSERT INTO search_export_settings(key,value,updated_at) VALUES('enabled',?,?)
               ON CONFLICT(key) DO UPDATE SET value=excluded.value,updated_at=excluded.updated_at""",
            ("1" if enabled else "0", now),
        )
    return bool(enabled)


def record_worker_heartbeat(worker_id: str, current_job_id: str = "") -> None:
    init_db()
    with _connect() as conn:
        conn.execute(
            """INSERT INTO search_export_worker_state(worker_id,heartbeat_at,current_job_id)
               VALUES(?,?,?) ON CONFLICT(worker_id) DO UPDATE SET
               heartbeat_at=excluded.heartbeat_at,current_job_id=excluded.current_job_id""",
            (str(worker_id)[:200], utc_text(), str(current_job_id)[:80]),
        )


def worker_available(max_age_seconds: int = WORKER_HEARTBEAT_MAX_AGE_SECONDS) -> bool:
    init_db()
    cutoff = utc_now() - timedelta(seconds=max(1, int(max_age_seconds)))
    with _connect() as conn:
        rows = conn.execute(
            "SELECT heartbeat_at FROM search_export_worker_state ORDER BY heartbeat_at DESC LIMIT 4"
        ).fetchall()
    return any((_parse_utc(row["heartbeat_at"]) or datetime.min.replace(tzinfo=timezone.utc)) >= cutoff for row in rows)


def resource_snapshot() -> ResourceSnapshot:
    available: int | None = None
    load_one: float | None = None
    try:
        load_one = float(os.getloadavg()[0])
    except (AttributeError, OSError):
        pass
    try:
        import psutil  # type: ignore

        available = int(psutil.virtual_memory().available)
        if load_one is None:
            load_one = float(psutil.getloadavg()[0])
    except (ImportError, AttributeError, OSError):
        if os.name != "nt":
            try:
                for line in Path("/proc/meminfo").read_text(encoding="ascii").splitlines():
                    if line.startswith("MemAvailable:"):
                        available = int(line.split()[1]) * 1024
                        break
            except (OSError, ValueError, IndexError):
                pass
    return ResourceSnapshot(available_bytes=available, load_one=load_one)


def resources_allow_start(snapshot: ResourceSnapshot | None = None) -> bool:
    current = snapshot or resource_snapshot()
    # Unknown telemetry is not proof that the explicit safety gate is met.
    if current.available_bytes is None or current.load_one is None:
        return False
    if current.available_bytes < 2 * 1024**3:
        return False
    if current.load_one > 4.0:
        return False
    return True


def request_fingerprint(payload: dict) -> str:
    stable = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return sha256(stable.encode("utf-8")).hexdigest()


def _decode_job(row: sqlite3.Row | dict | None) -> dict | None:
    if row is None:
        return None
    data = dict(row)
    try:
        data["scope"] = json.loads(str(data.pop("scope_json", "null")))
    except (TypeError, ValueError):
        data["scope"] = None
    try:
        data["manifest"] = json.loads(str(data.pop("manifest_json", "{}")))
    except (TypeError, ValueError):
        data["manifest"] = {}
    return data


def create_job(
    *,
    user_id: int,
    query_text: str,
    normalized_query: str,
    search_mode: str,
    scope: object,
    book_filter: str,
    citation_style: str,
    output_format: str,
    corpus_version: str,
    template_version: str,
    base_url: str,
    active_member: bool,
) -> tuple[dict, bool]:
    init_db()
    query_value = str(query_text or "")
    if (
        not query_value.strip()
        or len(query_value) > MAX_QUERY_CHARS
        or _ILLEGAL_XML_CONTROL_RE.search(query_value)
    ):
        raise SearchExportError(f"检索词不能为空、超过 {MAX_QUERY_CHARS} 个字符或含无效控制字符。")
    if search_mode not in VALID_MODES or output_format not in VALID_FORMATS:
        raise SearchExportError("导出参数无效。")
    if citation_style not in VALID_STYLES:
        raise SearchExportError("引用格式无效。")
    now = utc_now()
    now_text = utc_text(now)
    expiry_text = utc_text(now + timedelta(hours=RETENTION_HOURS))
    result_limit = MEMBER_RESULT_LIMIT if active_member else REGULAR_RESULT_LIMIT
    hourly_limit = MEMBER_HOURLY_LIMIT if active_member else REGULAR_HOURLY_LIMIT
    fingerprint_payload = {
        "query": normalized_query,
        "mode": search_mode,
        "scope": scope,
        "book": book_filter,
        "style": citation_style,
        "format": output_format,
        "corpus": corpus_version,
        "template": template_version,
        "content": EXPORT_CONTENT_VERSION,
        "base_url": base_url.rstrip("/"),
        "result_limit": result_limit,
    }
    fingerprint = request_fingerprint(fingerprint_payload)
    cutoff = utc_text(now - timedelta(hours=1))
    with _connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        reusable = conn.execute(
            """SELECT * FROM search_export_jobs
               WHERE user_id=? AND request_fingerprint=? AND expires_at>?
                 AND status IN ('queued','counting','collecting','rendering','packaging','complete')
               ORDER BY created_at DESC LIMIT 1""",
            (int(user_id), fingerprint, now_text),
        ).fetchone()
        if reusable is not None:
            candidate = _decode_job(reusable) or {}
            if candidate.get("status") != "complete" or completed_artifact(candidate) is not None:
                conn.commit()
                return candidate, True
            # The database must not advertise a vanished artifact as reusable.
            conn.execute(
                """UPDATE search_export_jobs SET status='expired',output_path='',output_name='',
                   output_size=0,updated_at=? WHERE id=?""",
                (now_text, str(candidate.get("id") or "")),
            )
        active = conn.execute(
            """SELECT COUNT(*) FROM search_export_jobs
               WHERE user_id=? AND status IN ('queued','counting','collecting','rendering','packaging')""",
            (int(user_id),),
        ).fetchone()[0]
        if int(active) >= MAX_ACTIVE_PER_USER:
            conn.rollback()
            raise QueueLimitError("你已有两个导出任务在排队或处理中，请完成后再试。")
        recent = conn.execute(
            "SELECT COUNT(*) FROM search_export_jobs WHERE user_id=? AND created_at>=?",
            (int(user_id), cutoff),
        ).fetchone()[0]
        if int(recent) >= hourly_limit:
            conn.rollback()
            raise RateLimitError(f"本小时最多创建 {hourly_limit} 个导出任务，请稍后再试。")
        job_id = uuid.uuid4().hex
        conn.execute(
            """INSERT INTO search_export_jobs(
               id,user_id,query_text,normalized_query,search_mode,scope_json,book_filter,
               citation_style,output_format,corpus_version,template_version,request_fingerprint,
               base_url,result_limit,status,created_at,updated_at,expires_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                job_id, int(user_id), query_text, normalized_query, search_mode,
                json.dumps(scope, ensure_ascii=False, sort_keys=True), book_filter,
                citation_style, output_format, corpus_version, template_version,
                fingerprint, base_url.rstrip("/"), result_limit, "queued",
                now_text, now_text, expiry_text,
            ),
        )
        row = conn.execute("SELECT * FROM search_export_jobs WHERE id=?", (job_id,)).fetchone()
        conn.commit()
    return _decode_job(row) or {}, False


def get_job(job_id: str, user_id: int | None = None) -> dict | None:
    init_db()
    query = "SELECT * FROM search_export_jobs WHERE id=?"
    params: tuple[object, ...] = (str(job_id),)
    if user_id is not None:
        query += " AND user_id=?"
        params += (int(user_id),)
    with _connect() as conn:
        return _decode_job(conn.execute(query, params).fetchone())


def update_job(job_id: str, **changes: object) -> dict | None:
    init_db()
    allowed = {
        "status", "result_count", "progress_done", "progress_total", "output_path",
        "output_name", "output_size", "manifest_json", "error", "lease_owner",
        "lease_expires_at", "expires_at",
    }
    values = {key: value for key, value in changes.items() if key in allowed}
    if "manifest_json" in values and not isinstance(values["manifest_json"], str):
        values["manifest_json"] = json.dumps(values["manifest_json"], ensure_ascii=False, sort_keys=True)
    values["updated_at"] = utc_text()
    assignments = ",".join(f"{key}=?" for key in values)
    with _connect() as conn:
        conn.execute(
            f"UPDATE search_export_jobs SET {assignments} WHERE id=?",
            (*values.values(), str(job_id)),
        )
        row = conn.execute("SELECT * FROM search_export_jobs WHERE id=?", (str(job_id),)).fetchone()
    return _decode_job(row)


def claim_next_job(worker_id: str, lease_seconds: int = MAX_RUNTIME_SECONDS + 60, *,
                   corpus_version: str | None = None,
                   template_version: str | None = None) -> dict | None:
    init_db()
    now = utc_now()
    now_text = utc_text(now)
    lease_text = utc_text(now + timedelta(seconds=max(60, int(lease_seconds))))
    version_clause = ""
    version_values: list[str] = []
    if corpus_version is not None:
        version_clause += " AND corpus_version=?"
        version_values.append(str(corpus_version))
    if template_version is not None:
        version_clause += " AND template_version=?"
        version_values.append(str(template_version))
    with _connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        # A killed worker may be replaced.  The first expired lease is returned
        # to the queue only while its two-attempt budget is not exhausted.
        conn.execute(
            """UPDATE search_export_jobs SET status='queued',lease_owner='',lease_expires_at=NULL,
               error='工作节点中断，任务已自动重试。',updated_at=?
               WHERE status IN ('counting','collecting','rendering','packaging')
                 AND lease_expires_at IS NOT NULL AND lease_expires_at<? AND attempts<?""",
            (now_text, now_text, MAX_ATTEMPTS),
        )
        conn.execute(
            """UPDATE search_export_jobs SET status='failed',lease_owner='',lease_expires_at=NULL,
               error='工作节点连续中断，任务已停止。',updated_at=?
               WHERE status IN ('counting','collecting','rendering','packaging')
                 AND lease_expires_at IS NOT NULL AND lease_expires_at<? AND attempts>=?""",
            (now_text, now_text, MAX_ATTEMPTS),
        )
        row = conn.execute(
            "SELECT id FROM search_export_jobs "
            "WHERE status='queued' AND expires_at>? AND attempts<? "
            + version_clause + " ORDER BY created_at LIMIT 1",
            (now_text, MAX_ATTEMPTS, *version_values),
        ).fetchone()
        if row is None:
            conn.commit()
            return None
        job_id = str(row["id"])
        changed = conn.execute(
            """UPDATE search_export_jobs SET status='counting',attempts=attempts+1,
               lease_owner=?,lease_expires_at=?,error='',updated_at=?
               WHERE id=? AND status='queued'""",
            (str(worker_id)[:200], lease_text, now_text, job_id),
        ).rowcount
        if not changed:
            conn.rollback()
            return None
        claimed = conn.execute("SELECT * FROM search_export_jobs WHERE id=?", (job_id,)).fetchone()
        conn.commit()
    return _decode_job(claimed)


def renew_lease(job_id: str, worker_id: str, lease_seconds: int = MAX_RUNTIME_SECONDS + 60) -> None:
    init_db()
    with _connect() as conn:
        conn.execute(
            """UPDATE search_export_jobs SET lease_expires_at=?,updated_at=?
               WHERE id=? AND lease_owner=?""",
            (
                utc_text(utc_now() + timedelta(seconds=max(60, int(lease_seconds)))),
                utc_text(), str(job_id), str(worker_id)[:200],
            ),
        )


def requeue_job(job_id: str, message: str) -> dict | None:
    job = get_job(job_id)
    if not job:
        return None
    if int(job.get("attempts") or 0) >= MAX_ATTEMPTS:
        return update_job(
            job_id, status="failed", error=str(message)[:500], lease_owner="", lease_expires_at=None
        )
    return update_job(
        job_id, status="queued", error=str(message)[:500], lease_owner="", lease_expires_at=None
    )


def pause_job(job_id: str, message: str) -> dict | None:
    """Release a claimed job without treating resource/citation preemption as a failure."""
    init_db()
    now_text = utc_text()
    with _connect() as conn:
        conn.execute(
            """UPDATE search_export_jobs
               SET status='queued',attempts=CASE WHEN attempts>0 THEN attempts-1 ELSE 0 END,
                   error=?,lease_owner='',lease_expires_at=NULL,updated_at=?
               WHERE id=? AND status IN ('counting','collecting','rendering','packaging')""",
            (str(message)[:500], now_text, str(job_id)),
        )
        row = conn.execute("SELECT * FROM search_export_jobs WHERE id=?", (str(job_id),)).fetchone()
    return _decode_job(row)


def _job_directory(job_id: str, create: bool = False) -> Path:
    root = ARTIFACT_ROOT.resolve()
    target = (root / str(job_id)).resolve()
    if target.parent != root:
        raise SearchExportError("导出目录无效。")
    if create:
        target.mkdir(parents=True, exist_ok=True, mode=0o700)
        _tighten_permissions(target, 0o700)
    return target


def completed_artifact(job: dict) -> Path | None:
    if str(job.get("status") or "") != "complete":
        return None
    job_id = str(job.get("id") or "")
    if not job_id:
        return None
    raw = str(job.get("output_path") or "")
    if not raw:
        return None
    try:
        job_directory = _job_directory(job_id).resolve()
    except SearchExportError:
        return None
    path = Path(raw).resolve()
    if path.parent != job_directory or not path.is_file():
        return None
    return path


def cleanup_expired() -> int:
    init_db()
    now_text = utc_text()
    with _connect() as conn:
        rows = conn.execute(
            "SELECT id FROM search_export_jobs WHERE expires_at<=? AND status!='expired'",
            (now_text,),
        ).fetchall()
    cleaned = 0
    for row in rows:
        job_id = str(row["id"])
        directory = _job_directory(job_id)
        if directory.exists():
            shutil.rmtree(directory)
        update_job(
            job_id, status="expired", output_path="", output_name="", output_size=0,
            lease_owner="", lease_expires_at=None,
        )
        cleaned += 1
    return cleaned


def _safe_filename(value: object, fallback: str = "引文检索") -> str:
    text = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", "_", str(value or "")).strip(" ._")
    text = re.sub(r"\s+", " ", text)
    return (text[:48] or fallback)


def _document_text(value: object) -> str:
    return _ILLEGAL_XML_CONTROL_RE.sub("", str(value or ""))


def _plain_context(value: object) -> str:
    return " ".join(_document_text(value).replace("\r", " ").replace("\n", " ").split())


def _context_segments(value: object) -> list[tuple[str, bool]]:
    text = _plain_context(value)
    parts = re.split(r"(\[\[H\]\]|\[\[/H\]\])", text)
    highlighted = False
    out: list[tuple[str, bool]] = []
    for part in parts:
        if part == "[[H]]":
            highlighted = True
        elif part == "[[/H]]":
            highlighted = False
        elif part:
            out.append((part, highlighted))
    return out


def _selected_citation(hit: dict, style: str) -> str:
    citations = hit.get("citations") if isinstance(hit.get("citations"), dict) else {}
    return _document_text(citations.get(style) or hit.get("citation") or "").strip()


def _title_for_hit(hit: dict) -> str:
    return _document_text(
        hit.get("book_title") or hit.get("book_short_title") or hit.get("display_title")
        or hit.get("book") or "资料"
    ).strip()


def _add_docx_hyperlink(paragraph, text: str, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "8F1D1D")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _set_docx_font(run, ascii_font: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    from docx.oxml.ns import qn

    run.font.name = ascii_font
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), ascii_font)
    fonts.set(qn("w:hAnsi"), ascii_font)
    fonts.set(qn("w:eastAsia"), east_asia)


def _configure_docx_styles(document) -> None:
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Inches, Pt, RGBColor

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    tokens = {
        "Normal": (11, "231D17", 0, 6, 1.25),
        "Heading 1": (16, "2E74B5", 18, 10, 1.0),
        "Heading 2": (13, "2E74B5", 14, 7, 1.0),
        "Heading 3": (12, "1F4D78", 10, 5, 1.0),
    }
    for name, (size, color, before, after, spacing) in tokens.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = spacing
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE if spacing != 1.0 else WD_LINE_SPACING.SINGLE


def _render_docx(path: Path, hits: list[dict], meta: dict) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    document = Document()
    _configure_docx_styles(document)
    document.core_properties.title = _document_text(meta["title"])
    document.core_properties.subject = "引文检索结果汇编"
    document.core_properties.author = "马克思主义文献检索"

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run(_document_text(meta["title"]))
    _set_docx_font(title_run)
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(0x8F, 0x1D, 0x1D)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sub_run = subtitle.add_run(
        f"检索词：{_document_text(meta['query'])}  ·  引用格式：{_document_text(meta['citation_label'])}  ·  "
        f"本册 {len(hits)} 条 / 共 {meta['total']} 条"
    )
    _set_docx_font(sub_run)
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(0x72, 0x67, 0x5D)

    current_book = None
    current_volume = None
    current_section = None
    for index, hit in enumerate(hits, start=int(meta.get("start_index") or 1)):
        book = _title_for_hit(hit)
        volume = int(hit.get("volume") or 0)
        section_title = _document_text(hit.get("section_title") or "（未识别篇章）").strip()
        if book != current_book:
            document.add_heading(book, level=1)
            current_book, current_volume, current_section = book, None, None
        if volume != current_volume:
            volume_label = _document_text(hit.get("display_title") or "").strip() or f"第 {volume} 卷"
            document.add_heading(volume_label, level=2)
            current_volume, current_section = volume, None
        if section_title != current_section:
            document.add_heading(section_title, level=3)
            current_section = section_title

        label = document.add_paragraph()
        label.paragraph_format.keep_with_next = True
        label_run = label.add_run(f"命中 {index}")
        _set_docx_font(label_run)
        label_run.bold = True
        label_run.font.color.rgb = RGBColor(0x65, 0x13, 0x13)

        context = document.add_paragraph()
        context.paragraph_format.keep_together = True
        for segment, highlighted in _context_segments(hit.get("context")):
            run = context.add_run(segment)
            _set_docx_font(run)
            if highlighted:
                run.bold = True
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                run.font.color.rgb = RGBColor(0x8F, 0x1D, 0x1D)

        citation = document.add_paragraph()
        citation.paragraph_format.space_before = Pt(4)
        citation.paragraph_format.space_after = Pt(8)
        citation_run = citation.add_run("出处：" + _selected_citation(hit, str(meta["citation_style"])))
        _set_docx_font(citation_run)
        citation_run.font.size = Pt(10)
        citation_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        link = str(hit.get("viewer_url") or "")
        if link.startswith(("https://", "http://")):
            spacer = citation.add_run("  ")
            _set_docx_font(spacer)
            _add_docx_hyperlink(citation, "打开原文页", link)

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("第 ")
    _set_docx_font(footer_run)
    footer_run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.extend((begin, instr, separate, placeholder, end))
    footer._p.append(field_run)
    footer_suffix = footer.add_run(" 页")
    _set_docx_font(footer_suffix)
    footer_suffix.font.size = Pt(9)
    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    document.save(path)


def _highlighted_html(value: object) -> str:
    return "".join(
        f"<mark>{html.escape(text)}</mark>" if highlighted else html.escape(text)
        for text, highlighted in _context_segments(value)
    )


def _render_html(path: Path, hits: list[dict], meta: dict) -> None:
    toc: list[str] = []
    body: list[str] = []
    toc_counter = 0
    current_book = None
    current_volume = None
    current_section = None
    for local_index, hit in enumerate(hits, start=1):
        absolute_index = int(meta.get("start_index") or 1) + local_index - 1
        book = _title_for_hit(hit)
        volume = int(hit.get("volume") or 0)
        section_title = _document_text(hit.get("section_title") or "（未识别篇章）").strip()
        if book != current_book:
            toc_counter += 1
            anchor = f"section-{toc_counter}"
            toc.append(f'<li class="toc-book"><a href="#{anchor}">{html.escape(book)}</a></li>')
            body.append(f'<h2 id="{anchor}">{html.escape(book)}</h2>')
            current_book, current_volume, current_section = book, None, None
        if volume != current_volume:
            volume_label = _document_text(hit.get("display_title") or "").strip() or f"第 {volume} 卷"
            toc_counter += 1
            anchor = f"section-{toc_counter}"
            toc.append(f'<li class="toc-volume"><a href="#{anchor}">{html.escape(volume_label)}</a></li>')
            body.append(f'<h3 id="{anchor}">{html.escape(volume_label)}</h3>')
            current_volume, current_section = volume, None
        if section_title != current_section:
            toc_counter += 1
            anchor = f"section-{toc_counter}"
            toc.append(f'<li class="toc-section"><a href="#{anchor}">{html.escape(section_title)}</a></li>')
            body.append(f'<h4 id="{anchor}">{html.escape(section_title)}</h4>')
            current_section = section_title
        citation = html.escape(_selected_citation(hit, str(meta["citation_style"])))
        link = str(hit.get("viewer_url") or "")
        link_html = (
            f'<a class="source" href="{html.escape(link, quote=True)}" target="_blank" rel="noopener">打开原文页</a>'
            if link.startswith(("https://", "http://")) else ""
        )
        body.append(
            f'<article class="hit"><div class="hit-no">命中 {absolute_index}</div>'
            f'<p class="context">{_highlighted_html(hit.get("context"))}</p>'
            f'<p class="citation">出处：{citation}</p>{link_html}</article>'
        )
    title = html.escape(_document_text(meta["title"]))
    query = html.escape(_document_text(meta["query"]))
    citation_label = html.escape(_document_text(meta["citation_label"]))
    document = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title><style>
:root{{--ink:#231d17;--muted:#72675d;--accent:#8f1d1d;--line:#decfbd;--paper:#fffdf9}}
*{{box-sizing:border-box}} body{{margin:0;background:#f4efe8;color:var(--ink);font:16px/1.75 "Microsoft YaHei","Noto Sans CJK SC",sans-serif}}
main{{max-width:980px;margin:0 auto;padding:36px 24px 80px;background:var(--paper);min-height:100vh}}
h1{{margin:0;color:var(--accent);font-family:"Noto Serif CJK SC","Songti SC",serif}} h2{{margin-top:40px;border-bottom:2px solid var(--line)}}
h3{{color:#2e74b5}} h4{{color:#1f4d78}} .meta{{color:var(--muted)}} nav{{padding:16px 20px;background:#f5eee6;border:1px solid var(--line);border-radius:12px}}
nav ul{{columns:2;padding-left:24px}} nav li{{break-inside:avoid}} .toc-volume{{margin-left:1em}} .toc-section{{margin-left:2em;color:var(--muted)}}
.hit{{margin:18px 0;padding:18px 20px;border:1px solid var(--line);border-radius:12px;break-inside:avoid}}
.hit-no{{font-weight:700;color:#651313}} .context{{font-family:"Noto Serif CJK SC","Songti SC",serif}} mark{{background:#fff0a8;color:#8f1d1d;font-weight:700;padding:0 .08em}}
.citation{{color:#555;font-size:.94rem}} .source{{display:inline-block;color:var(--accent);font-weight:700}} @media print{{body{{background:#fff}}main{{max-width:none;padding:0}}nav{{break-after:page}}.source{{word-break:break-all}}}}
@media(max-width:640px){{main{{padding:22px 16px}}nav ul{{columns:1}}}}
</style></head><body><main><h1>{title}</h1><p class="meta">检索词：{query} · 引用格式：{citation_label} · 本册 {len(hits)} 条 / 共 {int(meta['total'])} 条</p>
<nav><strong>目录</strong><ul>{''.join(toc)}</ul></nav>{''.join(body)}</main></body></html>"""
    path.write_text(document, encoding="utf-8")


def _write_part(path: Path, output_format: str, hits: list[dict], meta: dict) -> None:
    partial = path.with_name(path.name + ".part")
    partial.unlink(missing_ok=True)
    try:
        if output_format == "docx":
            _render_docx(partial, hits, meta)
        else:
            _render_html(partial, hits, meta)
        if not partial.is_file() or partial.stat().st_size <= 0:
            raise SearchExportError("导出文件生成失败。")
        os.replace(partial, path)
        _tighten_permissions(path, 0o600)
    finally:
        partial.unlink(missing_ok=True)


def _zip_parts(path: Path, parts: list[Path]) -> None:
    partial = path.with_name(path.name + ".part")
    partial.unlink(missing_ok=True)
    try:
        with zipfile.ZipFile(partial, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
            for part in parts:
                archive.write(part, arcname=part.name)
        os.replace(partial, path)
        _tighten_permissions(path, 0o600)
    finally:
        partial.unlink(missing_ok=True)


def _citation_label(style: str) -> str:
    return {
        "gb2025": "GB/T 7714—2025",
        "gb2015": "GB/T 7714—2015",
        "zgshkx": "《中国社会科学》",
        "mkszyj": "《马克思主义研究》",
    }.get(style, style)


def run_job(
    job_id: str,
    *,
    worker_id: str,
    count_hits: Callable[[dict, int], int],
    iter_hits: Callable[[dict, int], Iterator[dict]],
    continue_after_part: Callable[[], bool] | None = None,
) -> dict | None:
    """Execute one claimed job.  Exceptions are contained in this job lane."""
    import time

    job = get_job(job_id)
    if not job:
        return None
    started = time.monotonic()
    directory = _job_directory(job_id, create=True)
    for partial in directory.glob("*.part"):
        partial.unlink(missing_ok=True)
    try:
        if job.get("corpus_version") == "" or job.get("template_version") == "":
            raise SearchExportError("导出版本信息缺失。")
        result_limit = int(job.get("result_limit") or REGULAR_RESULT_LIMIT)
        count = int(count_hits(job, result_limit + 1))
        if count > result_limit:
            return update_job(
                job_id, status="too_large", result_count=count, progress_total=count,
                error=f"共检索到至少 {count} 条，超过当前账号 {result_limit} 条上限，请缩小检索范围。",
                lease_owner="", lease_expires_at=None,
            )
        if time.monotonic() - started > MAX_RUNTIME_SECONDS:
            raise RetryableExportError("任务计数超过 15 分钟，已暂停并重试。")
        update_job(
            job_id, status="collecting", result_count=count, progress_total=count,
            progress_done=0,
        )
        safe_query = _safe_filename(job.get("query_text"))
        extension = str(job["output_format"])
        title = f"“{job['query_text']}”引文检索汇编"
        base_meta = {
            "title": title,
            "query": str(job["query_text"]),
            "total": count,
            "citation_style": str(job["citation_style"]),
            "citation_label": _citation_label(str(job["citation_style"])),
        }
        manifest = job.get("manifest") if isinstance(job.get("manifest"), dict) else {}
        valid_existing = {
            str(item.get("name")): item
            for item in manifest.get("parts") or []
            if isinstance(item, dict) and str(item.get("name") or "")
        }
        completed_parts: list[dict] = []
        emitted = 0

        def finish_part(
            part_hits: list[dict], part_number: int, start_index: int,
            key: tuple[str, str, int],
        ) -> None:
            nonlocal completed_parts
            book_name = _safe_filename(key[1], "资料")
            volume_name = f"第{key[2]}卷" if key[2] else "资料"
            filename = f"{part_number:03d}_{book_name}_{volume_name}.{extension}"
            target = directory / filename
            prior = valid_existing.get(filename)
            if not (prior and target.is_file() and target.stat().st_size == int(prior.get("size") or -1)):
                update_job(job_id, status="rendering", progress_done=start_index - 1)
                _write_part(
                    target, extension, part_hits,
                    {**base_meta, "start_index": start_index},
                )
            size = target.stat().st_size
            if size > MAX_OUTPUT_BYTES:
                raise OutputTooLargeError("单个导出文件超过 128MB，任务已停止。")
            completed_parts.append({"name": filename, "size": size, "count": len(part_hits)})
            update_job(
                job_id,
                status="collecting",
                progress_done=start_index + len(part_hits) - 1,
                manifest_json={"parts": completed_parts},
            )
            renew_lease(job_id, worker_id)
            record_worker_heartbeat(worker_id, job_id)
            if continue_after_part is not None and not continue_after_part():
                raise PausedExportError("主站任务优先或系统资源紧张，导出已暂停并重新排队。")

        if count <= SINGLE_FILE_HIT_LIMIT:
            hits = list(iter_hits(job, result_limit + 1))
            emitted = len(hits)
            if emitted != count:
                raise RetryableExportError("检索数据在导出期间发生变化，任务将重新核对。")
            filename = f"{safe_query}_引文检索汇编.{extension}"
            target = directory / filename
            prior = valid_existing.get(filename)
            if not (prior and target.is_file() and target.stat().st_size == int(prior.get("size") or -1)):
                update_job(job_id, status="rendering")
                _write_part(target, extension, hits, {**base_meta, "start_index": 1})
            if target.stat().st_size > MAX_OUTPUT_BYTES:
                raise OutputTooLargeError("导出文件超过 128MB，任务已停止。")
            update_job(
                job_id,
                status="collecting",
                progress_done=count,
                manifest_json={
                    "parts": [{"name": filename, "size": target.stat().st_size, "count": count}]
                },
            )
            renew_lease(job_id, worker_id)
            record_worker_heartbeat(worker_id, job_id)
            if continue_after_part is not None and not continue_after_part():
                raise PausedExportError("主站任务优先或系统资源紧张，导出已暂停并重新排队。")
            final_path = target
            output_name = filename
        else:
            part_hits: list[dict] = []
            part_key: tuple[str, str, int] | None = None
            part_number = 0
            part_start = 1
            for hit in iter_hits(job, result_limit + 1):
                key = (
                    str(hit.get("book") or ""),
                    _title_for_hit(hit),
                    int(hit.get("volume") or 0),
                )
                if part_hits and (key != part_key or len(part_hits) >= PART_HIT_LIMIT):
                    part_number += 1
                    finish_part(part_hits, part_number, part_start, part_key or ("", "资料", 0))
                    part_start += len(part_hits)
                    part_hits = []
                if not part_hits:
                    part_key = key
                part_hits.append(hit)
                emitted += 1
                if emitted % 100 == 0 and time.monotonic() - started > MAX_RUNTIME_SECONDS:
                    raise RetryableExportError("任务运行超过 15 分钟，已暂停并重试。")
            if part_hits:
                part_number += 1
                finish_part(part_hits, part_number, part_start, part_key or ("", "资料", 0))
            if emitted != count:
                raise RetryableExportError("检索数据在导出期间发生变化，任务将重新核对。")
            update_job(job_id, status="packaging", progress_done=count)
            part_paths = [directory / str(item["name"]) for item in completed_parts]
            output_name = f"{safe_query}_引文检索汇编_{count}条.zip"
            final_path = directory / output_name
            _zip_parts(final_path, part_paths)
            if final_path.stat().st_size > MAX_OUTPUT_BYTES:
                raise OutputTooLargeError("压缩包超过 128MB，任务已停止。")
        completed = update_job(
            job_id,
            status="complete",
            progress_done=count,
            progress_total=count,
            output_path=str(final_path.resolve()),
            output_name=output_name,
            output_size=final_path.stat().st_size,
            expires_at=utc_text(utc_now() + timedelta(hours=RETENTION_HOURS)),
            error="",
            lease_owner="",
            lease_expires_at=None,
        )
        return completed
    except OutputTooLargeError as exc:
        return update_job(
            job_id, status="failed", error=str(exc)[:500], lease_owner="", lease_expires_at=None
        )
    except PausedExportError as exc:
        return pause_job(job_id, str(exc))
    except RetryableExportError as exc:
        return requeue_job(job_id, str(exc))
    except Exception as exc:
        return requeue_job(job_id, f"导出失败：{str(exc)[:430]}")


def public_payload(job: dict) -> dict:
    return {
        "id": str(job.get("id") or ""),
        "status": str(job.get("status") or ""),
        "format": str(job.get("output_format") or ""),
        "mode": str(job.get("search_mode") or ""),
        "query": str(job.get("query_text") or ""),
        "result_count": int(job.get("result_count") or 0),
        "result_limit": int(job.get("result_limit") or 0),
        "progress_done": int(job.get("progress_done") or 0),
        "progress_total": int(job.get("progress_total") or 0),
        "output_name": str(job.get("output_name") or ""),
        "output_size": int(job.get("output_size") or 0),
        "error": str(job.get("error") or ""),
        "created_at": str(job.get("created_at") or ""),
        "updated_at": str(job.get("updated_at") or ""),
        "expires_at": str(job.get("expires_at") or ""),
    }
