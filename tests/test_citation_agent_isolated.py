from __future__ import annotations

import ast
import hashlib
import json
import tarfile
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement

import citation_agent_bridge as bridge
import citation_agent_queue as queue
import citation_agent_test_backend as test_tasks
import citation_assistant as public_tasks
from scripts.build_citation_agent_release import build as build_agent_release
from scripts import citation_agent_worker as network_worker


REAL_ARTICLES = {
    "wang": (
        Path(r"C:\Users\10108\Downloads\王贤生：纪检监察制度的百年演进与新时代中国特色反腐败共同体的建构路径（0918版）.docx"),
        "d8344ed60549f9060736de8e99825992ff4fc7ae8aecbcf9da158630cdcd8ac2",
    ),
    "pan": (
        Path(r"C:\Users\10108\Downloads\潘云龙、翟云：用户数据何由平台做主？——数据领域的劳动赋权论误用批判.docx"),
        "1aa9ade78959ac2c1bf865f5390211881f02c3a082dc1cb7cebb357205f6dd92",
    ),
    "technology": (
        Path(r"D:\360MoveData\Users\10108\Documents\xwechat_files\wxid_97bddsw2dytg22_8760\msg\file\2026-09\从法律赋权到技术赋权——数智时代西方左翼的一个社会主义道路转向.docx"),
        "3eea11c236703787962e5a9fc340930c98a905787a39833a48a025cbefba9613",
    ),
}


@pytest.fixture()
def isolated_test_store(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    db_path = tmp_path / "agent-test.sqlite3"
    artifact_root = tmp_path / "agent-test-artifacts"
    monkeypatch.setattr(test_tasks.core, "DB_PATH", db_path)
    monkeypatch.setattr(test_tasks.core, "ARTIFACT_ROOT", artifact_root)
    monkeypatch.setattr(test_tasks, "DB_PATH", db_path)
    monkeypatch.setattr(test_tasks, "ARTIFACT_ROOT", artifact_root)
    test_tasks.init_db()
    return tmp_path


def _docx(path: Path) -> bytes:
    document = Document()
    document.add_heading("正文", level=1)
    document.add_paragraph("马克思指出：“社会生活在本质上是实践的。”")
    document.save(path)
    return path.read_bytes()


def _job(store: Path) -> dict:
    return test_tasks.create_job(
        7, "匿名样本.docx", _docx(store / "paper.docx"),
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )


def test_admin_test_requires_explicit_supported_scope(isolated_test_store: Path) -> None:
    with pytest.raises(test_tasks.CitationAssistantError, match="请先指定"):
        test_tasks.create_job(
            7, "匿名样本.docx", _docx(isolated_test_store / "no-scope.docx"),
            recognition_depth="direct_only", scope_tokens=[],
            corpus_sha256="corpus", template_version="template",
        )

    job = _job(isolated_test_store)
    test_tasks.run_extraction(str(job["id"]))
    current = test_tasks.get_job(str(job["id"]), 7)
    section_ids = [str(item["id"]) for item in current.get("sections") or []]
    with pytest.raises(test_tasks.CitationAssistantError, match="请先指定"):
        test_tasks.set_analysis_config(
            str(job["id"]), 7, section_ids=section_ids, scope_tokens=[],
        )


def test_public_and_agent_test_stores_are_distinct() -> None:
    assert public_tasks is not test_tasks.core
    assert public_tasks.DB_PATH != test_tasks.core.DB_PATH
    assert public_tasks.ARTIFACT_ROOT != test_tasks.core.ARTIFACT_ROOT


def test_agent_test_jobs_remain_owner_isolated(isolated_test_store: Path) -> None:
    job = _job(isolated_test_store)
    assert test_tasks.get_job(str(job["id"]), 7)
    assert test_tasks.get_job(str(job["id"]), 8) is None
    assert test_tasks.delete_job(str(job["id"]), 8) is False
    assert test_tasks.get_job(str(job["id"]), 7)


def _queue_test_job_for_analysis(
    store: Path,
    *,
    corpus_sha256: str = "corpus",
    template_version: str = "template",
) -> dict:
    job = test_tasks.create_job(
        7, "匿名恢复样本.docx", _docx(store / f"recover-{template_version}.docx"),
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256=corpus_sha256, template_version=template_version,
    )
    test_tasks.run_extraction(str(job["id"]))
    current = test_tasks.get_job(str(job["id"]), 7)
    sections = [str(item["id"]) for item in current.get("sections") or []]
    return test_tasks.set_analysis_config(
        str(job["id"]), 7, section_ids=sections, scope_tokens=["book:文集"],
    )


def test_version_recovery_only_requeues_safe_current_runtime_jobs(
    isolated_test_store: Path,
) -> None:
    analysis_error = "引文模板已变更，请重新创建测试任务。"
    export_error = "引文模板已变更，为保证结果可复现，已拒绝导出。"

    recoverable = _queue_test_job_for_analysis(isolated_test_store)
    test_tasks.update_job(
        str(recoverable["id"]), status="failed", error=analysis_error,
        progress_done=9, progress_total=12, lease_owner="stale", lease_expires_at="2099-01-01T00:00:00+00:00",
    )

    wrong_template = _queue_test_job_for_analysis(
        isolated_test_store, template_version="old-template",
    )
    test_tasks.update_job(str(wrong_template["id"]), status="failed", error=analysis_error)

    wrong_corpus = _queue_test_job_for_analysis(
        isolated_test_store, corpus_sha256="old-corpus",
    )
    test_tasks.update_job(str(wrong_corpus["id"]), status="failed", error=analysis_error)

    missing_file = _queue_test_job_for_analysis(isolated_test_store)
    Path(str(missing_file["extraction_path"])).unlink()
    test_tasks.update_job(str(missing_file["id"]), status="failed", error=analysis_error)

    missing_sections = _queue_test_job_for_analysis(isolated_test_store)
    test_tasks.update_job(str(missing_sections["id"]), status="failed", error=analysis_error)
    with test_tasks.core._connect() as conn:
        conn.execute(
            "UPDATE citation_assistant_jobs SET selected_sections_json='[]' WHERE id=?",
            (str(missing_sections["id"]),),
        )

    unrelated = _queue_test_job_for_analysis(isolated_test_store)
    test_tasks.update_job(str(unrelated["id"]), status="failed", error="模型暂不可用")

    expired = _queue_test_job_for_analysis(isolated_test_store)
    test_tasks.update_job(str(expired["id"]), status="failed", error=analysis_error)
    with test_tasks.core._connect() as conn:
        conn.execute(
            "UPDATE citation_assistant_jobs SET expires_at=? WHERE id=?",
            ("2000-01-01T00:00:00+00:00", str(expired["id"])),
        )

    exportable = _job(isolated_test_store)
    test_tasks.update_job(
        str(exportable["id"]), status="failed", error=export_error, candidate_count=1,
    )
    empty_export = _job(isolated_test_store)
    test_tasks.update_job(str(empty_export["id"]), status="failed", error=export_error)

    recovered = test_tasks.recover_jobs_for_loaded_runtime(
        "corpus",
        "template",
        analysis_errors=(analysis_error,),
        export_errors=(export_error,),
    )

    assert recovered == {"analysis": 1, "export": 1}
    analysis = test_tasks.get_job(str(recoverable["id"]), 7)
    assert analysis["status"] == "queued"
    assert analysis["error"] == ""
    assert analysis["progress_done"] == 0
    assert analysis["progress_total"] == 0
    assert analysis["lease_owner"] == ""
    assert analysis["selected_sections"] == recoverable["selected_sections"]
    assert test_tasks.get_job(str(exportable["id"]), 7)["status"] == "exporting"
    for job in (
        wrong_template, wrong_corpus, missing_file, missing_sections,
        unrelated, expired, empty_export,
    ):
        assert test_tasks.get_job(str(job["id"]), 7)["status"] == "failed"


def test_version_filtered_claim_leaves_stale_runtime_jobs_unclaimed(
    isolated_test_store: Path,
) -> None:
    stale = test_tasks.create_job(
        7, "旧版本.docx", _docx(isolated_test_store / "stale.docx"),
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="old-template",
    )
    current = _job(isolated_test_store)

    claimed = test_tasks.claim_next_job(
        "worker:current", lease_seconds=60,
        corpus_sha256="corpus", template_version="template",
    )

    assert claimed and claimed["id"] == current["id"]
    assert test_tasks.get_job(str(stale["id"]), 7)["status"] == "extracting"
    assert test_tasks.get_job(str(stale["id"]), 7)["lease_owner"] == ""


def test_queue_rejects_original_document_metadata() -> None:
    with pytest.raises(queue.AgentQueueError, match="禁止字段"):
        queue.validate_request({
            "round": 1, "recognition_depth": "direct_only",
            "allowed_public_books": ["文集"], "personal_source_ids": [],
            "records": [{"record_id": "r1", "kind": "quote", "fragment": "社会生活在本质上是实践的", "filename": "论文.docx"}],
        })


def test_queue_round_trip_keeps_only_validated_json(tmp_path: Path) -> None:
    path = tmp_path / "queue" / "queue.sqlite3"
    request = {
        "round": 1, "recognition_depth": "direct_only", "prompt_version": "v1",
        "allowed_public_books": ["文集"], "personal_source_ids": ["p1"],
        "records": [{"record_id": "r1", "kind": "quote", "fragment": "社会生活在本质上是实践的"}],
        "observations": [],
    }
    task_id = queue.enqueue(request, path=path)
    claimed = queue.claim("agent:test", path=path)
    assert claimed and claimed["id"] == task_id
    queue.complete(task_id, "agent:test", {
        "actions": [{
            "action_id": "a1", "record_id": "r1", "tool": "exact_fragment",
            "fragment": "社会生活在本质上是实践的", "keywords": [],
        }],
    }, path=path)
    result = queue.result(task_id, path=path)
    assert result and result["status"] == "complete"
    assert result["response"]["actions"][0]["tool"] == "exact_fragment"


def test_second_round_observation_keeps_attempted_query() -> None:
    request = queue.validate_request({
        "round": 2, "planning_effort": "fast",
        "recognition_depth": "direct_only", "prompt_version": "v2",
        "allowed_public_books": ["文集"], "personal_source_ids": [],
        "records": [{"record_id": "r1", "kind": "quote", "fragment": "社会生活在本质上是实践的"}],
        "observations": [{
            "record_id": "r1", "action_id": "a1", "tool": "exact_fragment",
            "fragment": "社会生活在本质上", "keywords": [], "hit_count": 0,
            "exact_count": 0, "unique_source_count": 0, "result": "none",
        }],
    })
    observation = request["observations"][0]
    assert observation["tool"] == "exact_fragment"
    assert observation["fragment"] == "社会生活在本质上"
    assert "不要重复" in network_worker._messages(request)[0]["content"]


def test_queue_uses_same_normalized_owner_for_unicode_hostname(tmp_path: Path) -> None:
    path = tmp_path / "queue-unicode-owner" / "queue.sqlite3"
    request = {
        "round": 1, "recognition_depth": "direct_only", "prompt_version": "v1",
        "allowed_public_books": ["文集"], "personal_source_ids": [],
        "records": [{"record_id": "r1", "kind": "quote", "fragment": "社会生活在本质上是实践的"}],
        "observations": [],
    }
    task_id = queue.enqueue(request, path=path)
    worker_id = "agent:本机测试:42"
    assert queue.claim(worker_id, path=path)["id"] == task_id
    queue.complete(task_id, worker_id, {"actions": []}, path=path)
    assert queue.result(task_id, path=path)["status"] == "complete"


def test_direct_mode_rejects_paraphrase_action() -> None:
    request = queue.validate_request({
        "round": 1, "recognition_depth": "direct_only",
        "allowed_public_books": ["文集"], "personal_source_ids": [],
        "records": [{"record_id": "r1", "kind": "direct", "fragment": "劳动创造价值并形成社会关系"}],
    })
    with pytest.raises(queue.AgentQueueError, match="禁止观点转述"):
        queue.validate_response({
            "actions": [{
                "action_id": "a1", "record_id": "r1", "tool": "keyword_cooccurrence",
                "fragment": "", "keywords": ["劳动", "价值"],
            }],
        }, request)


def test_personal_scope_is_opaque_to_network_agent() -> None:
    corpus = SimpleNamespace(books={"文集": []})
    scope, public, personal = bridge._scope(
        corpus, ["book:mylib:42", "book:文集"],
    )
    assert scope == {"文集": None}
    assert public == ["文集"]
    assert personal == ["p1"]
    assert "42" not in "".join(personal)


def test_unresolved_candidates_are_read_only_and_not_actionable(isolated_test_store: Path) -> None:
    job = _job(isolated_test_store)
    test_tasks.replace_candidates(str(job["id"]), [{
        "kind": "audit", "paper_text": "无法确认的句子", "existing_note_text": "《未知著作》第1页",
        "issue_code": "unverifiable", "issue_label": "疑似属于所选文库，但尚未核实",
        "review_bucket": "unresolved", "evidence_origin": "agent", "source_options": [],
    }])
    current = test_tasks.get_job(str(job["id"]), 7)
    assert current and current["candidate_count"] == 0
    assert current["unresolved_count"] == 1
    unresolved = test_tasks.list_candidates(str(job["id"]), bucket="unresolved")
    candidate_id = unresolved["items"][0]["id"]
    with pytest.raises(test_tasks.CitationAssistantError, match="只读"):
        test_tasks.save_decisions(str(job["id"]), 7, [{"id": candidate_id, "decision": "accepted"}])


def test_viewpoint_support_is_counted_separately_and_never_bulk_accepted(
    isolated_test_store: Path,
) -> None:
    job = _job(isolated_test_store)
    assert job["viewpoint_suggestion_count"] == 0
    test_tasks.replace_candidates(str(job["id"]), [
        {
            "paper_text": "直接引文候选", "issue_code": "suggest_add",
            "review_bucket": "actionable", "evidence_origin": "agent",
            "evidence_level": "near_text", "source_options": [],
        },
        {
            "paper_text": "观点依据候选", "issue_code": "suggest_add",
            "review_bucket": "actionable", "evidence_origin": "agent",
            "evidence_level": "paraphrase", "source_options": [],
        },
    ])
    test_tasks.update_job(
        str(job["id"]), agent_verified_count=1, viewpoint_suggestion_count=1,
    )
    result = test_tasks.bulk_decide_pending(str(job["id"]), 7, "accepted")
    assert result["updated"] == 1
    rows = test_tasks.all_candidates(str(job["id"]), bucket="actionable")
    decisions = {item["paper_text"]: item["decision"] for item in rows}
    assert decisions == {"直接引文候选": "accepted", "观点依据候选": "pending"}
    current = test_tasks.get_job(str(job["id"]), 7)
    assert current and current["agent_verified_count"] == 1
    assert current["viewpoint_suggestion_count"] == 1


def test_agent_run_persists_direct_and_viewpoint_counts_separately(
    isolated_test_store: Path,
) -> None:
    job = _job(isolated_test_store)
    test_tasks.save_agent_run(str(job["id"]), {
        "status": "complete", "mode": "admin_live", "model": "test-model",
        "prompt_version": "test-prompt", "verified_match_count": 3,
        "direct_verified_match_count": 2, "viewpoint_suggestion_count": 1,
    })
    row = test_tasks.get_agent_run(str(job["id"]))
    assert row and row["verified_match_count"] == 3
    assert row["direct_verified_match_count"] == 2
    assert row["viewpoint_suggestion_count"] == 1


def test_export_receives_actionable_candidates_only(
    isolated_test_store: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = _job(isolated_test_store)
    test_tasks.replace_candidates(str(job["id"]), [
        {"paper_text": "可操作", "issue_code": "suggest_add", "review_bucket": "actionable", "source_options": []},
        {"paper_text": "只读", "issue_code": "unverifiable", "review_bucket": "unresolved", "source_options": []},
    ])
    test_tasks.update_job(str(job["id"]), status="review_ready")
    received: list[list[dict]] = []
    pdf_options: list[dict] = []

    def fake_pdf(_job: dict, candidates: list[dict], **kwargs) -> Path:
        received.append(candidates)
        pdf_options.append(kwargs)
        return isolated_test_store / "out.pdf"

    def fake_docx(_job: dict, candidates: list[dict]) -> Path:
        received.append(candidates)
        return isolated_test_store / "out.docx"

    monkeypatch.setattr(test_tasks.core, "export_pdf", fake_pdf)
    monkeypatch.setattr(test_tasks.core, "export_docx", fake_docx)
    test_tasks.run_export(str(job["id"]))
    assert len(received) == 2
    assert all([item["paper_text"] for item in group] == ["可操作"] for group in received)
    assert pdf_options == [{
        "source_docx": str(isolated_test_store / "out.docx"),
        "suppress_agent_comments_for_conversion": True,
    }]


def test_agent_failure_keeps_deterministic_result(
    isolated_test_store: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = _job(isolated_test_store)
    test_tasks.run_extraction(str(job["id"]))
    extracted_job = test_tasks.get_job(str(job["id"]), 7)
    assert extracted_job and extracted_job["status"] == "awaiting_sections"
    section_ids = [str(item["id"]) for item in extracted_job.get("sections") or []]
    test_tasks.set_analysis_config(
        str(job["id"]), 7, section_ids=section_ids, scope_tokens=["book:文集"],
    )
    option = _Hit().to_dict()
    monkeypatch.setattr(test_tasks.core, "_phrase_records", lambda *_args, **_kwargs: ([], {}))
    monkeypatch.setattr(test_tasks.core, "analyze_extraction", lambda *_args, **_kwargs: ([{
        "kind": "generate", "section_id": section_ids[0], "paragraph_index": 1,
        "raw_start": 0, "raw_end": 10, "paper_text": "确定性结果",
        "match_type": "exact", "score": 100, "fuzzy_errors": 0,
        "issue_code": "suggest_add", "issue_label": "建议补注",
        "source_options": [option], "selected_option": 0,
        "proposed_citation": "测试引文，第12页。", "auto_selected": True,
    }], "mkszyj", 1.0))

    def broken_agent(**_kwargs):
        raise RuntimeError("agent unavailable")

    test_tasks.run_analysis(
        str(job["id"]), SimpleNamespace(books={"文集": []}), agent_callback=broken_agent,
    )
    finished = test_tasks.get_job(str(job["id"]), 7)
    assert finished and finished["status"] == "review_ready"
    assert finished["agent_status"] == "degraded"
    assert finished["candidate_count"] == 1
    assert test_tasks.list_candidates(str(job["id"]), bucket="actionable")["items"][0]["paper_text"] == "确定性结果"


class _Hit:
    def to_dict(self) -> dict:
        return {
            "book": "文集", "volume": 1, "source_file": "wenji-1.pdf",
            "display_title": "《马克思恩格斯文集》第1卷", "pdf_pages": [25],
            "printed_pages": [12], "context": "社会生活在本质上是实践的。",
            "citation": "测试引文，第12页。", "citations": {"mkszyj": "测试引文，第12页。"},
            "match_type": "exact", "score": 100, "fuzzy_errors": 0,
        }


class _Corpus:
    def locate_quote(self, text: str, **kwargs):
        assert kwargs["book_scope"] == {"文集": None}
        return [_Hit()] if text == "社会生活在本质上是实践的" else []


def test_auxiliary_index_and_note_pages_are_not_source_evidence(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class AuxiliaryHit:
        def __init__(self, section: str):
            self.section = section

    monkeypatch.setattr(bridge.tasks.core, "_option_from_hit", lambda hit: {
        "book": "文集", "volume": 5, "source_file": f"{hit.section}.pdf",
        "printed_pages": ["1097"], "section_title": hit.section,
    })
    monkeypatch.setattr(
        bridge.tasks.core, "_option_signature",
        lambda option: (option["source_file"], tuple(option["printed_pages"])),
    )
    options = bridge._options([
        AuxiliaryHit("名目索引"), AuxiliaryHit("注释"), AuxiliaryHit("资本的生产过程"),
    ])
    assert [item["section_title"] for item in options] == ["资本的生产过程"]


def test_local_executor_can_auto_adopt_only_complete_unique_exact() -> None:
    record = {
        "_raw": "社会生活在本质上是实践的", "raw_text": "社会生活在本质上是实践的",
        "paragraph_index": 1, "section_id": "body", "raw_start": 3, "raw_end": 16,
    }
    action = {
        "action_id": "a1", "record_id": "r1", "tool": "exact_fragment",
        "fragment": "社会生活在本质上是实践的", "keywords": [],
    }
    observation, candidate = bridge._execute_action(
        action, record, _Corpus(), {"文集": None}, "mkszyj", "direct_only",
    )
    assert observation["result"] == "unique_exact"
    assert candidate and candidate["evidence_level"] == "exact"
    assert candidate["auto_selected"] is True

    action["fragment"] = "本质上是实践的"
    observation, candidate = bridge._execute_action(
        action, record, _Corpus(), {"文集": None}, "mkszyj", "direct_only",
    )
    assert observation["result"] == "none"
    assert candidate is None


def test_rejected_local_action_returns_safe_second_round_observation() -> None:
    record = {
        "_raw": "社会生活在本质上是实践的", "raw_text": "社会生活在本质上是实践的",
        "paragraph_index": 1, "section_id": "body", "raw_start": 3, "raw_end": 16,
    }
    observation, candidate = bridge._execute_action({
        "action_id": "a1", "record_id": "r1", "tool": "exact_fragment",
        "fragment": "这不是输入中的片段", "keywords": [],
    }, record, _Corpus(), {"文集": None}, "mkszyj", "direct_only")
    assert candidate is None and observation["result"] == "rejected"
    assert observation["tool"] == "exact_fragment"
    assert observation["fragment"] == "这不是输入中的片段"


def test_broad_paraphrase_evidence_is_suppressed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    options = [
        {"book": "文集", "volume": index, "source_file": f"v{index}.pdf",
         "printed_pages": [str(index)], "pdf_pages": [index], "match_type": "fuzzy",
         "score": 88, "citation": f"出处{index}", "citations": {"mkszyj": f"出处{index}"}}
        for index in range(1, 5)
    ]
    monkeypatch.setattr(bridge, "_options", lambda _hits: options)
    corpus = SimpleNamespace(keyword_cooccurrence=lambda *_args, **_kwargs: [object()] * 4)
    record = {
        "_raw": "历史发展需要理论和实践彼此统一", "raw_text": "历史发展需要理论和实践彼此统一",
        "_kind": "paraphrase",
        "paragraph_index": 1, "section_id": "body", "raw_start": 0, "raw_end": 18,
    }
    observation, candidate = bridge._execute_action({
        "action_id": "a1", "record_id": "r1", "tool": "keyword_cooccurrence",
        "fragment": "", "keywords": ["历史发展", "理论", "实践"],
    }, record, corpus, {"文集": None}, "mkszyj", "direct_and_paraphrase")
    assert observation["result"] == "ambiguous"
    assert candidate is None


def test_paraphrase_requires_three_local_concepts_and_all_concepts_in_window() -> None:
    calls: list[dict] = []

    class Corpus:
        def keyword_cooccurrence(self, keywords, **kwargs):
            calls.append({"keywords": keywords, **kwargs})
            return []

    record = {
        "_raw": "劳动过程体现人与自然之间的物质变换关系",
        "raw_text": "劳动过程体现人与自然之间的物质变换关系",
        "_kind": "paraphrase",
        "paragraph_index": 1, "section_id": "body", "raw_start": 0, "raw_end": 20,
    }
    rejected, candidate = bridge._execute_action({
        "action_id": "a1", "record_id": "r1", "tool": "keyword_cooccurrence",
        "fragment": "", "keywords": ["劳动过程", "自然"],
    }, record, Corpus(), {"文集": None}, "mkszyj", "direct_and_paraphrase")
    assert rejected["result"] == "rejected" and candidate is None and calls == []

    observation, candidate = bridge._execute_action({
        "action_id": "a2", "record_id": "r1", "tool": "keyword_cooccurrence",
        "fragment": "", "keywords": ["劳动过程", "自然", "物质变换"],
    }, record, Corpus(), {"文集": None}, "mkszyj", "direct_and_paraphrase")
    assert observation["result"] == "none" and candidate is None
    assert calls == [{
        "keywords": ["劳动过程", "自然", "物质变换"],
        "window": 180, "min_distinct": 3, "occ_cap": 120,
        "book_scope": {"文集": None}, "expand_synonyms": False,
    }]


def test_paraphrase_lane_uses_unverified_note_locator_not_ordinary_prose() -> None:
    record = {
        "raw_text": "哲学不能止于解释世界而应致力于改变世界",
        "paragraph_index": 0, "section_id": "body", "raw_start": 0, "raw_end": 20,
        "quoted": False,
    }
    redacted, record_map, deferred = bridge._redacted_records(
        [record], [], "direct_and_paraphrase",
    )
    assert redacted == [] and record_map == {} and deferred == 0

    locator = {
        "kind": "audit", "review_bucket": "unresolved",
        "verification_scope": "locator_only", "paragraph_index": 0,
        "section_id": "body", "raw_start": 0, "raw_end": 20,
        "paper_text": "哲学不能止于解释世界而应致力于改变世界",
    }
    redacted, record_map, deferred = bridge._redacted_records(
        [record], [locator], "direct_and_paraphrase",
    )
    assert deferred == 0 and len(redacted) == 1
    assert redacted[0]["kind"] == "paraphrase"
    local_record = record_map[redacted[0]["record_id"]]
    assert not bridge._action_allowed({
        "tool": "near_quote", "fragment": "解释世界而应致力于改变世界", "keywords": [],
    }, local_record, "direct_and_paraphrase")
    assert bridge._action_allowed({
        "tool": "keyword_cooccurrence", "fragment": "",
        "keywords": ["哲学", "解释世界", "改变世界"],
    }, local_record, "direct_and_paraphrase")

    direct, direct_map, direct_deferred = bridge._redacted_records(
        [record], [locator], "direct_only",
    )
    assert direct == [] and direct_map == {} and direct_deferred == 0


def test_agent_source_link_keeps_all_highlights_and_safe_return_contract() -> None:
    script = (Path(__file__).resolve().parents[1] / "static" / "citation-agent-test" / "citation-agent-test.js").read_text(
        encoding="utf-8",
    )
    viewer = (Path(__file__).resolve().parents[1] / "templates" / "viewer.html").read_text(
        encoding="utf-8",
    )
    assert "context.matchAll" in script
    assert "matches.join(' ')" in script
    assert "highlight.slice(0,320)" in script
    assert "params.set('return_to',location.pathname+location.search)" in script
    assert 'class="ca-source-link" target="_blank"' not in script
    assert "params.set('as',previewIdentity)" in script
    assert "function withPreviewIdentity(url)" in script
    assert "withPreviewIdentity(job.docx_url)" in script
    assert "withPreviewIdentity(item.page_url)" in script
    # 最新主站阅读器不接收测试通道的返回链接；不得为了
    # 管理员隔离页面反向修改阅读器核心行为。
    assert "viewer_return_url" not in viewer


def test_admin_ui_names_viewpoint_discovery_as_a_separate_function() -> None:
    root = Path(__file__).resolve().parents[1]
    template = (root / "templates" / "citation_agent_test.html").read_text(encoding="utf-8")
    script = (root / "static" / "citation-agent-test" / "citation-agent-test.js").read_text(
        encoding="utf-8",
    )
    assert "观点依据发现（实验）" in template
    assert "不证明论文语句出自该文献" in template
    assert "一键采信直接引文" in template
    assert "插注＋校注（推荐）" in template
    assert "新增上标、标号和注文统一使用浅蓝色" in template
    assert "const modeLabels={both:'插注＋校注',generate:'插注',audit:'校注'}" in script
    assert "endnote:'浅蓝色尾注插入'" in script
    assert "这不是引文命中" in script
    assert "catViewpointSuggestions" in script


def test_external_or_unselected_notes_do_not_enter_regular_review() -> None:
    configs = [
        SimpleNamespace(key="文集", title="《马克思恩格斯文集》", short_title="《文集》", citation_title="马克思恩格斯文集"),
        SimpleNamespace(key="全集", title="《马克思恩格斯全集》", short_title="《全集》", citation_title="马克思恩格斯全集"),
    ]
    corpus = SimpleNamespace(
        books={"文集": [], "全集": []}, book_configs=configs,
        locate_quote=lambda *_args, **_kwargs: [],
    )
    base = {
        "kind": "audit", "section_id": "body", "paragraph_index": 1,
        "raw_start": 0, "raw_end": 20, "paper_text": "这是一段需要核验但当前没有命中的引文文字",
        "issue_code": "unverifiable", "issue_label": "本站无法核验", "source_options": [],
        "auto_selected": False,
    }
    kept, counts = test_tasks._partition_deterministic([
        {**base, "existing_note_text": "《站外学术专著》，某出版社，第20页。"},
        {**base, "paragraph_index": 2, "existing_note_text": "《马克思恩格斯全集》第1卷，第20页。"},
    ], corpus, ["book:文集"], "mkszyj")
    assert kept == []
    assert counts == {"out_of_scope_count": 1, "skipped_no_evidence_count": 1}


def test_selected_library_unverified_note_is_read_only() -> None:
    config = SimpleNamespace(
        key="文集", title="《马克思恩格斯文集》", short_title="《文集》",
        citation_title="马克思恩格斯文集",
    )
    corpus = SimpleNamespace(
        books={"文集": []}, book_configs=[config], locate_quote=lambda *_args, **_kwargs: [],
    )
    kept, counts = test_tasks._partition_deterministic([{
        "kind": "audit", "section_id": "body", "paragraph_index": 1,
        "raw_start": 0, "raw_end": 20, "paper_text": "这是一段需要核验但当前没有命中的引文文字",
        "existing_note_text": "《马克思恩格斯文集》第1卷，第20页。",
        "issue_code": "unverifiable", "issue_label": "本站无法核验",
        "source_options": [], "auto_selected": False,
    }], corpus, ["book:文集"], "mkszyj")
    assert counts == {"out_of_scope_count": 0, "skipped_no_evidence_count": 0}
    assert len(kept) == 1 and kept[0]["review_bucket"] == "unresolved"


@pytest.mark.parametrize(
    "issue", ["wrong_page", "wrong_source", "wrong_version", "format_mismatch"],
)
def test_admin_test_auto_adopts_safe_exact_correction_inside_full_quote(issue: str) -> None:
    quote = "社会生活在本质上是实践的。"
    paragraph_text = f"马克思指出：\u201c{quote}\u201d"
    start = paragraph_text.index(quote)
    candidate = {
        "kind": "audit", "section_id": "body", "paragraph_index": 0,
        "raw_start": start, "raw_end": start + len(quote), "paper_text": quote,
        "existing_note_text": "《马克思恩格斯文集》第1卷，第99页。",
        "match_type": "exact", "score": 100, "fuzzy_errors": 0,
        "issue_code": issue, "issue_label": issue,
        "source_options": [_Hit().to_dict()], "auto_selected": False,
    }
    corpus = SimpleNamespace(books={"文集": []})
    kept, counts = test_tasks._partition_deterministic(
        [candidate], corpus, ["book:文集"], "mkszyj",
        extracted={"paragraphs": [{"index": 0, "text": paragraph_text}]},
    )
    assert counts == {"out_of_scope_count": 0, "skipped_no_evidence_count": 0}
    assert len(kept) == 1
    assert kept[0]["verification_scope"] == "full_quote"
    assert kept[0]["evidence_level"] == "exact"
    assert kept[0]["review_bucket"] == "actionable"
    assert kept[0]["auto_selected"] is True


def test_field_paragraph_is_searchable_but_forced_into_read_only_bucket() -> None:
    quote = "国要有国法，党要有党规党法。"
    paragraph_text = f"邓小平指出：“{quote}”"
    paragraph = {
        "index": 0, "section_id": "body", "text": paragraph_text,
        "unsupported": True, "has_field": True, "tracked": False,
    }
    records, _ = public_tasks._phrase_records([paragraph], {"body"})
    assert records == []
    records, _ = public_tasks._phrase_records(
        [paragraph], {"body"}, include_readonly_structures=True,
    )
    assert len(records) == 1
    assert records[0]["quoted"] is True and records[0]["write_blocked"] is True

    start = paragraph_text.index(quote)
    candidate = {
        "kind": "generate", "section_id": "body", "paragraph_index": 0,
        "raw_start": start, "raw_end": start + len(quote), "paper_text": quote,
        "match_type": "exact", "score": 100, "fuzzy_errors": 0,
        "issue_code": "suggest_add", "issue_label": "建议新增注释",
        "source_options": [_Hit().to_dict()], "auto_selected": True,
    }
    kept, counts = test_tasks._partition_deterministic(
        [candidate], SimpleNamespace(books={"文集": []}), ["book:文集"], "mkszyj",
        extracted={"paragraphs": [paragraph]},
    )
    assert counts == {"out_of_scope_count": 0, "skipped_no_evidence_count": 0}
    assert kept[0]["review_bucket"] == "unresolved"
    assert kept[0]["verification_scope"] == "readonly_structure"
    assert kept[0]["evidence_level"] == "exact"
    assert kept[0]["auto_selected"] is False

    audit_candidate = {
        **candidate, "kind": "audit", "existing_note_kind": "footnote",
        "existing_note_id": 1, "existing_note_text": "原脚注",
    }
    paragraph["safe_comment_spans"] = [[0, len(paragraph_text)]]
    kept, _counts = test_tasks._partition_deterministic(
        [audit_candidate], SimpleNamespace(books={"文集": []}), ["book:文集"], "mkszyj",
        extracted={"paragraphs": [paragraph]},
    )
    assert kept[0]["review_bucket"] == "actionable"
    assert kept[0]["writeback_mode"] == "comment"


def test_plain_reference_heading_forces_tail_out_of_default_sections(tmp_path: Path) -> None:
    path = tmp_path / "reference-tail.docx"
    document = Document()
    document.add_heading("五、结语", level=1)
    document.add_paragraph("正文仍需核验：“社会生活在本质上是实践的。”")
    document.add_paragraph("参考文献")
    document.add_paragraph("［1］某书中也出现“不应被当作论文正文的参考文献引句”。")
    document.add_paragraph("Abstract")
    document.save(path)

    extracted = public_tasks.extract_docx(path)
    reference_sections = [
        item for item in extracted["sections"] if "参考文献" in str(item["title"])
    ]
    assert len(reference_sections) == 1
    assert reference_sections[0]["default_selected"] is False
    selected = {
        str(item["id"]) for item in extracted["sections"] if item.get("default_selected")
    }
    records, _ = public_tasks._phrase_records(extracted["paragraphs"], selected)
    assert [record["raw_text"] for record in records] == ["社会生活在本质上是实践的。"]


def test_agent_test_scope_widget_restores_initial_tokens() -> None:
    script = (
        Path(__file__).resolve().parents[1] / "static" / "book-scope" / "book-scope.js"
    ).read_text(encoding="utf-8")
    assert "initialTokens" in script
    assert "fromTokens(opts.initialTokens)" in script


def test_pdf_failure_keeps_successful_word_export(
    isolated_test_store: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = _job(isolated_test_store)
    test_tasks.update_job(str(job["id"]), status="review_ready")
    calls: list[str] = []
    word_path = isolated_test_store / "word-only.docx"
    word_path.write_bytes(b"docx")

    def fake_docx(_job: dict, _candidates: list[dict]) -> Path:
        calls.append("docx")
        return word_path

    def fake_pdf(_job: dict, _candidates: list[dict], **_kwargs) -> Path:
        calls.append("pdf")
        raise test_tasks.CitationAssistantError("PDF 转换器未就绪")

    monkeypatch.setattr(test_tasks.core, "export_docx", fake_docx)
    monkeypatch.setattr(test_tasks.core, "export_pdf", fake_pdf)
    test_tasks.run_export(str(job["id"]))
    current = test_tasks.get_job(str(job["id"]), 7)
    assert calls == ["docx", "pdf"]
    assert current and current["status"] == "complete"
    assert current["output_docx_path"] == str(word_path)
    assert current["output_pdf_path"] == ""
    assert "Word 已生成" in current["error"]


def test_pdf_converter_uses_private_profile_and_cleans_it(
    isolated_test_store: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = isolated_test_store / "probe.docx"
    _docx(source)
    output = isolated_test_store / "probe.pdf"
    monkeypatch.setenv("CITATION_ASSISTANT_SOFFICE", str(isolated_test_store / "soffice.exe"))
    (isolated_test_store / "soffice.exe").write_bytes(b"binary")
    seen: dict[str, object] = {}

    def fake_run(command: list[str], **kwargs) -> SimpleNamespace:
        seen["command"] = command
        seen["env"] = kwargs["env"]
        target = Path(command[command.index("--outdir") + 1]) / "probe.pdf"
        target.write_bytes(b"%PDF-1.7\n" + b"x" * 1200)
        return SimpleNamespace(returncode=0, stdout="converted")

    monkeypatch.setattr(public_tasks.subprocess, "run", fake_run)
    public_tasks._convert_annotated_docx_to_pdf(source, output)
    assert output.stat().st_size > 1000
    assert str((seen["env"] or {}).get("HOME", "")).startswith(str(isolated_test_store))
    assert not list(isolated_test_store.glob(".lo-profile-*"))
    assert not list(isolated_test_store.glob(".lo-output-*"))


def test_invalid_pdf_timeout_falls_back_to_180(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CITATION_ASSISTANT_PDF_TIMEOUT", "invalid")
    assert public_tasks._pdf_timeout_seconds() == 180


def test_pdf_converter_reports_missing_component(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("CITATION_ASSISTANT_SOFFICE", raising=False)
    monkeypatch.setattr(public_tasks.shutil, "which", lambda _command: None)
    with pytest.raises(public_tasks.CitationAssistantError, match="转换组件"):
        public_tasks._convert_annotated_docx_to_pdf(Path("missing.docx"), Path("missing.pdf"))


@pytest.mark.parametrize("failure", ["nonzero", "empty", "timeout"])
def test_pdf_converter_failure_cleans_private_directories(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, failure: str,
) -> None:
    source = tmp_path / "failure.docx"
    _docx(source)
    office = tmp_path / "soffice.exe"
    office.write_bytes(b"binary")
    monkeypatch.setenv("CITATION_ASSISTANT_SOFFICE", str(office))

    def fake_run(command: list[str], **_kwargs):
        if failure == "timeout":
            raise public_tasks.subprocess.TimeoutExpired(command, 30)
        if failure == "empty":
            (Path(command[command.index("--outdir") + 1]) / "failure.pdf").write_bytes(b"")
            return SimpleNamespace(returncode=0, stdout="")
        return SimpleNamespace(returncode=4, stdout="conversion failed")

    monkeypatch.setattr(public_tasks.subprocess, "run", fake_run)
    with pytest.raises(public_tasks.CitationAssistantError):
        public_tasks._convert_annotated_docx_to_pdf(source, tmp_path / "failure.pdf")
    assert not list(tmp_path.glob(".lo-profile-*"))
    assert not list(tmp_path.glob(".lo-output-*"))


def test_pdf_retry_reuses_final_word_without_reexporting(
    isolated_test_store: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = _job(isolated_test_store)
    final_word = Path(str(job["input_path"]))
    test_tasks.update_job(
        str(job["id"]), status="complete", output_docx_path=str(final_word),
        word_export_status="ready", pdf_export_status="failed",
    )
    queued = test_tasks.queue_pdf_retry(str(job["id"]), 7)
    assert queued["status"] == "exporting"
    assert queued["pdf_export_status"] == "converting"
    with pytest.raises(test_tasks.CitationAssistantError):
        test_tasks.queue_pdf_retry(str(job["id"]), 7)

    output = isolated_test_store / "retried.pdf"
    monkeypatch.setattr(
        test_tasks.core, "export_docx",
        lambda *_args, **_kwargs: pytest.fail("PDF retry must not regenerate Word"),
    )

    def fake_pdf(
        _job: dict, _candidates: list[dict], *, source_docx: Path,
        suppress_agent_comments_for_conversion: bool = False,
    ) -> Path:
        assert Path(source_docx) == final_word
        assert suppress_agent_comments_for_conversion is True
        output.write_bytes(b"%PDF-retried")
        return output

    monkeypatch.setattr(test_tasks.core, "export_pdf", fake_pdf)
    test_tasks.run_export(str(job["id"]))
    current = test_tasks.get_job(str(job["id"]), 7)
    assert current and current["status"] == "complete"
    assert current["word_export_status"] == "ready"
    assert current["pdf_export_status"] == "ready"
    assert current["output_pdf_path"] == str(output)


def test_admin_test_keeps_unquoted_exact_wording_for_manual_review() -> None:
    sentence = "社会生活在本质上是实践的。"
    candidate = {
        "kind": "audit", "section_id": "body", "paragraph_index": 0,
        "raw_start": 0, "raw_end": len(sentence), "paper_text": sentence,
        "existing_note_text": "《马克思恩格斯文集》第1卷，第99页。",
        "match_type": "exact", "score": 100, "fuzzy_errors": 0,
        "issue_code": "wrong_page", "issue_label": "页码错误",
        "source_options": [_Hit().to_dict()], "auto_selected": False,
    }
    kept, _counts = test_tasks._partition_deterministic(
        [candidate], SimpleNamespace(books={"文集": []}), ["book:文集"], "mkszyj",
        extracted={"paragraphs": [{"index": 0, "text": sentence}]},
    )
    assert kept[0]["verification_scope"] == "note_sentence"
    assert kept[0]["evidence_level"] == "exact"
    assert kept[0]["review_bucket"] == "actionable"
    assert kept[0]["auto_selected"] is False


def test_legacy_pending_exact_quote_is_backfilled_without_overwriting_human_decision(
    isolated_test_store: Path,
) -> None:
    job = _job(isolated_test_store)
    quote = "社会生活在本质上是实践的。"
    paragraph_text = f"马克思指出：\u201c{quote}\u201d"
    start = paragraph_text.index(quote)
    extraction_path = isolated_test_store / "legacy-extraction.json"
    extraction_path.write_text(
        json.dumps({
            "paragraphs": [{"index": 0, "text": paragraph_text}],
        }, ensure_ascii=False),
        encoding="utf-8",
    )
    test_tasks.update_job(str(job["id"]), extraction_path=str(extraction_path))
    base = {
        "kind": "audit", "section_id": "body", "paragraph_index": 0,
        "raw_start": start, "raw_end": start + len(quote), "paper_text": quote,
        "existing_note_text": "《马克思恩格斯文集》第1卷，第99页。",
        "match_type": "exact", "score": 100, "fuzzy_errors": 0,
        "issue_code": "wrong_page", "issue_label": "页码错误",
        "source_options": [_Hit().to_dict()], "auto_selected": False,
    }
    test_tasks.replace_candidates(str(job["id"]), [base, {**base, "paragraph_index": 1}])
    rows = test_tasks.all_candidates(str(job["id"]), bucket="actionable")
    with test_tasks.core._connect() as conn:
        conn.execute(
            "UPDATE citation_assistant_candidates SET decision='rejected',auto_selected=0 WHERE id=?",
            (int(rows[1]["id"]),),
        )
        test_tasks._backfill_existing_candidate_rules(conn)
    rows = test_tasks.all_candidates(str(job["id"]), bucket="actionable")
    assert rows[0]["verification_scope"] == "full_quote"
    assert rows[0]["decision"] == "accepted" and rows[0]["auto_selected"] is True
    assert rows[1]["decision"] == "rejected" and rows[1]["auto_selected"] is False


def test_locator_only_page_is_read_only_and_not_near_text(
    isolated_test_store: Path,
) -> None:
    job = _job(isolated_test_store)
    locator = {
        "kind": "audit", "section_id": "body", "paragraph_index": 0,
        "raw_start": 0, "raw_end": 16, "paper_text": "论文作者的概括性表述",
        "existing_note_text": "《马克思恩格斯文集》第1卷，第12页。",
        "match_type": "locator", "score": 0, "fuzzy_errors": None,
        "issue_code": "unverifiable", "issue_label": "本站无法核验",
        "source_options": [_Hit().to_dict()], "auto_selected": False,
        "review_bucket": "actionable", "evidence_level": "near_text",
    }
    test_tasks.replace_candidates(str(job["id"]), [locator])
    current = test_tasks.get_job(str(job["id"]), 7)
    assert current and current["candidate_count"] == 0
    assert current["unresolved_count"] == 1
    item = test_tasks.list_candidates(str(job["id"]), bucket="unresolved")["items"][0]
    assert item["verification_scope"] == "locator_only"
    assert item["evidence_level"] == ""
    assert item["decision"] == "informational"
    assert item["issue_label"] == "仅定位到原注所写页码，文字未核验"
    with pytest.raises(test_tasks.CitationAssistantError, match="只读"):
        test_tasks.save_decisions(
            str(job["id"]), 7, [{"id": item["id"], "decision": "accepted"}],
        )


def test_locator_is_manual_viewpoint_only_in_opt_in_mode() -> None:
    locator = {
        "kind": "audit", "paragraph_index": 1, "raw_start": 0, "raw_end": 18,
        "paper_text": "这是依据原著内容形成的一段概括性表述",
        "existing_note_text": "《马克思恩格斯文集》第2卷，第546页。",
        "match_type": "locator", "review_bucket": "unresolved",
        "evidence_level": "", "verification_scope": "locator_only",
        "source_options": [_Hit().to_dict()], "auto_selected": False,
    }
    direct = test_tasks._promote_locator_viewpoints(
        [locator], {}, "direct_only",
    )[0]
    assert direct["review_bucket"] == "unresolved"

    viewpoint = test_tasks._promote_locator_viewpoints(
        [locator], {}, "direct_and_paraphrase",
    )[0]
    assert viewpoint["review_bucket"] == "actionable"
    assert viewpoint["text_match_level"] == "paraphrase"
    assert viewpoint["source_resolution"] == "locator_only"
    assert viewpoint["writeback_mode"] == "none"
    assert viewpoint["auto_selected"] is False
    assert {"paraphrase_never_auto", "locator_without_text_match"}.issubset(
        viewpoint["reason_codes"],
    )


def test_agent_additions_cannot_cross_insertion_and_proofreading_modes() -> None:
    additions = [
        {"kind": "generate", "evidence_level": "exact"},
        {"kind": "audit", "evidence_level": "exact"},
        {"kind": "audit", "evidence_level": "paraphrase"},
    ]
    assert [item["evidence_level"] for item in test_tasks._filter_mode_additions(
        additions, "generate",
    )] == ["exact", "paraphrase"]
    assert [item["evidence_level"] for item in test_tasks._filter_mode_additions(
        additions, "audit",
    )] == ["exact", "paraphrase"]
    assert test_tasks._filter_mode_additions(additions, "both") == additions


def test_combined_mode_proofreading_hit_does_not_suppress_insertion_lane() -> None:
    record = {
        "paragraph_index": 3, "raw_start": 5, "raw_end": 20,
        "raw_text": "社会生活在本质上是实践的", "context_text": "马克思指出社会生活在本质上是实践的",
        "quoted": True,
    }
    audit = {
        "kind": "audit", "paragraph_index": 3, "raw_start": 5, "raw_end": 20,
        "review_bucket": "actionable", "source_options": [{"printed_pages": [1]}],
    }
    combined, _record_map, _deferred = bridge._redacted_records(
        [record], [audit], "direct_only", "both",
    )
    assert len(combined) == 1 and combined[0]["kind"] == "quote"
    audit_only, _record_map, _deferred = bridge._redacted_records(
        [record], [audit], "direct_only", "audit",
    )
    assert audit_only == []


def test_admin_ui_distinguishes_locator_from_textual_evidence() -> None:
    script = (
        Path(__file__).resolve().parents[1] / "static" / "citation-agent-test" /
        "citation-agent-test.js"
    ).read_text(encoding="utf-8")
    assert "原注页码定位" in script
    assert "文字未核验" in script
    assert "引号内完整文字" in script
    assert "locatorOnly?'':" in script
    assert "逐字核对一致" in script
    assert "校注意见只写入 Word 批注，不改正文" in script
    assert ".map(reasonLabel)" in script
    assert "class=\"ca-citation-edit cat-citation-edit\"" in script
    assert "scrollHeight" in script
    assert "legacyIssueLabels" in script
    assert "esc(issueLabel(candidate))" in script


def test_network_worker_does_not_import_site_or_corpus() -> None:
    path = Path(__file__).resolve().parents[1] / "scripts" / "citation_agent_worker.py"
    tree = ast.parse(path.read_text(encoding="utf-8"))
    imported = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported.add(node.module)
    assert "app" not in imported
    assert "citation_assistant" not in imported
    assert "citation_agent_test_backend" not in imported
    assert "search" not in imported


def test_network_worker_initializes_queue_before_first_cleanup(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    events: list[str] = []
    monkeypatch.setenv("CITATION_AGENT_TEST_MODE", "admin_live")
    monkeypatch.setattr(
        network_worker.queue,
        "init_queue",
        lambda path: events.append("init"),
    )
    monkeypatch.setattr(
        network_worker.queue,
        "cleanup",
        lambda: events.append("cleanup"),
    )

    def no_task(*_args, **_kwargs):
        events.append("claim")
        return None

    monkeypatch.setattr(network_worker.queue, "claim", no_task)
    assert network_worker.run(once=True, poll_seconds=0) == 0
    assert events == ["init", "cleanup", "claim"]


def test_deterministic_test_worker_does_not_import_web_application() -> None:
    path = Path(__file__).resolve().parents[1] / "scripts" / "citation_agent_test_worker.py"
    source = path.read_text(encoding="utf-8")
    tree = ast.parse(source)
    imported = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported.add(node.module)
    assert "app" not in imported
    assert "recover_jobs_for_loaded_runtime" in source
    assert "recover_pdf_position_failures" in source
    assert "corpus_sha256=loaded_corpus_sha256" in source
    assert "template_version=loaded_template_version" in source
    assert "runtime.template_version() != loaded_template_version" in source
    assert "signal.signal(signal.SIGTERM, _request_stop)" in source


def test_agent_release_bundle_is_strictly_allowlisted(tmp_path: Path) -> None:
    bundle = build_agent_release(tmp_path / "citation-agent.tar.gz")
    with tarfile.open(bundle, "r:gz") as archive:
        names = set(archive.getnames())
    assert names == {
        "citation_agent_queue.py", "scripts/citation_agent_worker.py", "SHA256SUMS",
    }


def test_network_worker_requires_proxy_and_allowlisted_https(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("HTTPS_PROXY", raising=False)
    monkeypatch.delenv("https_proxy", raising=False)
    monkeypatch.setenv("CITATION_AGENT_REQUIRE_PROXY", "1")
    with pytest.raises(network_worker.AgentWorkerError, match="出口代理"):
        network_worker._opener()

    monkeypatch.setenv("CITATION_AGENT_API_KEY", "x" * 20)
    monkeypatch.setenv("CITATION_AGENT_ALLOWED_HOSTS", "api.deepseek.com")
    monkeypatch.setenv("CITATION_AGENT_API_URL", "http://api.deepseek.com")
    with pytest.raises(network_worker.AgentWorkerError, match="HTTPS"):
        network_worker._endpoint()
    monkeypatch.setenv("CITATION_AGENT_API_URL", "https://example.com")
    with pytest.raises(network_worker.AgentWorkerError, match="允许名单"):
        network_worker._endpoint()


def test_network_worker_defaults_to_flash_thinking(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CITATION_AGENT_API_KEY", "x" * 20)
    monkeypatch.setenv("CITATION_AGENT_ALLOWED_HOSTS", "api.deepseek.com")
    monkeypatch.setenv("CITATION_AGENT_API_URL", "https://api.deepseek.com")
    monkeypatch.delenv("CITATION_AGENT_MODEL", raising=False)
    monkeypatch.delenv("CITATION_AGENT_THINKING", raising=False)
    assert network_worker._endpoint()[2] == "deepseek-v4-flash"
    assert network_worker._thinking_mode() == "enabled"
    assert network_worker._thinking_mode({"planning_effort": "fast"}) == "disabled"
    assert network_worker._thinking_mode({"planning_effort": "deep"}, repair=True) == "disabled"
    monkeypatch.setenv("CITATION_AGENT_THINKING", "disabled")
    assert network_worker._thinking_mode() == "disabled"


def test_invalid_agent_batch_does_not_cancel_other_batches(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payloads: dict[str, dict] = {}
    calls: list[str] = []

    def fake_enqueue(payload: dict, **_kwargs) -> str:
        task_id = f"q{len(payloads) + 1}"
        payloads[task_id] = payload
        return task_id

    def fake_wait(task_id: str, **_kwargs) -> dict:
        calls.append(task_id)
        if len(calls) == 1:
            return {"status": "failed", "error_code": "invalid_action_schema"}
        request = payloads[task_id]
        item = request["records"][0]
        return {"status": "complete", "response": {"actions": [{
            "action_id": f"a{len(calls)}", "record_id": item["record_id"],
            "tool": "exact_fragment", "fragment": item["fragment"], "keywords": [],
        }]}}

    monkeypatch.setattr(bridge.agent_queue, "enqueue", fake_enqueue)
    monkeypatch.setattr(bridge.agent_queue, "wait_result", fake_wait)
    monkeypatch.setattr(bridge.agent_queue, "delete", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(bridge.tasks, "update_job", lambda *_args, **_kwargs: None)
    records = [{
        "paragraph_index": index, "section_id": "body", "raw_start": 0, "raw_end": 15,
        "raw_text": "社会生活在本质上是实践的", "quoted": True,
    } for index in range(1, 10)]
    additions, report = bridge.run_agent(
        job={"id": "test", "recognition_depth": "direct_only", "scope": ["book:文集"]},
        extracted={}, records=records, deterministic_candidates=[], resolved_style="mkszyj",
        corpus=SimpleNamespace(books={"文集": []}, locate_quote=_Corpus().locate_quote),
    )
    assert len(calls) >= 2
    assert report["status"] == "degraded"
    assert report["verified_match_count"] >= 1
    assert any(item.get("evidence_origin") == "agent" for item in additions)


@pytest.mark.parametrize("name", ["wang", "pan", "technology"])
def test_three_article_golden_files_are_hash_pinned(name: str) -> None:
    path, expected = REAL_ARTICLES[name]
    if not path.is_file():
        pytest.skip(f"real golden article is not mounted: {path}")
    assert hashlib.sha256(path.read_bytes()).hexdigest() == expected


def test_wang_reference_fields_are_read_without_mutating_the_source() -> None:
    path, _expected = REAL_ARTICLES["wang"]
    if not path.is_file():
        pytest.skip("Wang golden article is not mounted")
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    extracted = public_tasks.extract_docx(path)
    fields = [
        note for paragraph in extracted["paragraphs"] for note in paragraph.get("note_refs", [])
        if note.get("kind") == "reference_field"
    ]
    assert extracted["flags"]["reference_field_count"] == 77
    assert extracted["flags"]["reference_field_occurrence_count"] == len(fields) == 112
    assert all(note.get("readonly") and note.get("text") for note in fields)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_pan_short_footnote_quotes_are_complete_anchors() -> None:
    path, _expected = REAL_ARTICLES["pan"]
    if not path.is_file():
        pytest.skip("Pan golden article is not mounted")
    extracted = public_tasks.extract_docx(path)
    anchors = {}
    for paragraph in extracted["paragraphs"]:
        for note in paragraph.get("note_refs", []):
            if note.get("id") in {29, 53}:
                anchors[int(note["id"])] = public_tasks._note_anchor(
                    paragraph["text"], int(note["offset"]),
                )[2]
    assert anchors == {29: "盗窃他人的劳动时间", 53: "自由的有意识的活动"}


def test_technology_article_body_styles_do_not_create_fake_sections() -> None:
    path, _expected = REAL_ARTICLES["technology"]
    if not path.is_file():
        pytest.skip("technology golden article is not mounted")
    extracted = public_tasks.extract_docx(path)
    sections = extracted["sections"]
    assert len(sections) == 4
    assert sections[0]["title"] == "正文导语"
    assert [section["start"] for section in sections] == [0, 8, 26, 45]
    assert all(section["default_selected"] for section in sections)


def test_external_quote_without_selected_source_cue_is_count_only() -> None:
    records = [{
        "paragraph_index": 1, "section_id": "body", "raw_start": 3, "raw_end": 20,
        "raw_text": "这是某位外部作者的直接引文", "context_text": "洛克写道：“这是某位外部作者的直接引文。”",
        "quoted": True,
    }]
    redacted, record_map, deferred = bridge._redacted_records(records, [], "direct_only")
    assert redacted == []
    assert record_map == {}
    assert deferred == 0


def test_candidate_api_exposes_review_contract_fields(isolated_test_store: Path) -> None:
    job = _job(isolated_test_store)
    test_tasks.replace_candidates(str(job["id"]), [{
        "kind": "generate", "paper_text": "社会生活在本质上是实践的",
        "match_type": "exact", "score": 100, "fuzzy_errors": 0,
        "issue_code": "suggest_add", "issue_label": "建议新增注释",
        "review_bucket": "actionable", "evidence_level": "exact",
        "verification_scope": "full_quote", "source_options": [_Hit().to_dict()],
        "auto_selected": True,
    }])
    item = test_tasks.all_candidates(str(job["id"]))[0]
    assert item["text_match_level"] == "exact"
    assert item["source_resolution"] == "unique"
    assert item["writeback_mode"] == "footnote"
    assert "exact_text" in item["reason_codes"]


def test_agent_run_records_terminal_accounting(isolated_test_store: Path) -> None:
    job = _job(isolated_test_store)
    test_tasks.save_agent_run(str(job["id"]), {
        "status": "complete", "mode": "admin_live", "direct_status": "complete",
        "paraphrase_status": "complete", "eligible_record_count": 5,
        "processed_record_count": 5, "deferred_record_count": 0,
        "completion_reason": "all_records_terminal",
    })
    row = test_tasks.get_agent_run(str(job["id"]))
    assert row and row["direct_status"] == "complete"
    assert row["paraphrase_status"] == "complete"
    assert row["eligible_record_count"] == row["processed_record_count"] == 5
    assert row["deferred_record_count"] == 0


def test_confirmed_field_paragraph_writes_true_word_comment(
    isolated_test_store: Path,
) -> None:
    path = isolated_test_store / "field-paper.docx"
    document = Document()
    paragraph = document.add_paragraph()
    prefix = "马克思指出：“社会生活在本质上是实践的。”"
    paragraph.add_run(prefix)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(f"{{{public_tasks.W_NS}}}fldCharType", "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = ' HYPERLINK \\l "Ref_1" '
    separate = OxmlElement("w:fldChar")
    separate.set(f"{{{public_tasks.W_NS}}}fldCharType", "separate")
    result = OxmlElement("w:t")
    result.text = "[1]"
    end_field = OxmlElement("w:fldChar")
    end_field.set(f"{{{public_tasks.W_NS}}}fldCharType", "end")
    for node in (begin, instruction, separate, result, end_field):
        run._r.append(node)
    document.save(path)
    job = test_tasks.create_job(
        7, path.name, path.read_bytes(), recognition_depth="direct_only",
        scope_tokens=["book:文集"], corpus_sha256="corpus", template_version="template",
    )
    current = test_tasks.get_job(str(job["id"]), 7)
    start = prefix.index("社会生活")
    finish = start + len("社会生活在本质上是实践的")
    output = public_tasks.export_docx(current, [{
        "kind": "audit", "paragraph_index": 0, "raw_start": start, "raw_end": finish,
        "paper_text": "社会生活在本质上是实践的", "decision": "accepted",
        "writeback_mode": "comment", "text_match_level": "exact",
        "source_resolution": "unique", "reason_codes": ["proofreading_comment_only"],
        "issue_label": "已核验原文；采信后将作为 Word 批注写入，不改动正文",
        "proposed_citation": "马克思：《关于费尔巴哈的提纲》，第1页。",
        "source_options": [{"display_title": "关于费尔巴哈的提纲", "printed_pages": [1]}],
    }])
    with zipfile.ZipFile(output) as package:
        names = set(package.namelist())
        assert "word/comments.xml" in names
        body = public_tasks._parse_xml(package.read("word/document.xml"))
        comments = public_tasks._parse_xml(package.read("word/comments.xml"))
        assert len(body.xpath(".//w:commentRangeStart", namespaces=public_tasks.NS)) == 1
        assert len(body.xpath(".//w:commentRangeEnd", namespaces=public_tasks.NS)) == 1
        assert len(body.xpath(".//w:commentReference", namespaces=public_tasks.NS)) == 1
        assert len(comments.xpath("./w:comment", namespaces=public_tasks.NS)) == 1
        instruction_text = "".join(body.xpath(".//w:instrText/text()", namespaces=public_tasks.NS))
        assert 'HYPERLINK \\l "Ref_1"' in instruction_text
        assert public_tasks._paragraph_text(body.xpath("./w:body//w:p", namespaces=public_tasks.NS)[0]) == prefix + "[1]"
        comment_text = "".join(comments.xpath("./w:comment//w:t/text()", namespaces=public_tasks.NS))
        assert "审核原因：校注意见只写入 Word 批注，不改正文" in comment_text
        assert "proofreading_comment_only" not in comment_text


def _accepted_exact_candidate(*, kind: str, paragraph_index: int = 1) -> dict:
    quote = "社会生活在本质上是实践的。"
    paragraph = "马克思指出：“社会生活在本质上是实践的。”"
    start = paragraph.index(quote)
    return {
        "kind": kind, "paragraph_index": paragraph_index,
        "raw_start": start, "raw_end": start + len(quote), "paper_text": quote,
        "decision": "accepted", "auto_selected": kind == "generate",
        "writeback_mode": "footnote" if kind == "generate" else "comment",
        "text_match_level": "exact", "source_resolution": "unique",
        "match_type": "exact", "score": 100, "fuzzy_errors": 0,
        "issue_code": "suggest_add" if kind == "generate" else "wrong_page",
        "issue_label": "建议插注" if kind == "generate" else "原脚注页码需要校注",
        "existing_note_kind": "footnote" if kind == "audit" else "",
        "existing_note_id": 1 if kind == "audit" else None,
        "existing_note_text": "《马克思恩格斯文集》第1卷，第99页。" if kind == "audit" else "",
        "proposed_citation": "马克思：《关于费尔巴哈的提纲》，第1页。",
        "source_options": [{"display_title": "关于费尔巴哈的提纲", "printed_pages": [1]}],
    }


def test_comment_uses_verified_note_markers_and_never_internal_ids() -> None:
    candidate = _accepted_exact_candidate(kind="audit")
    candidate.update({
        "existing_note_kind": "footnote", "existing_note_id": 66,
        "reason_codes": ["exact_text", "proofreading_comment_only"],
    })
    comment = public_tasks._candidate_comment_text(candidate)
    assert "对应脚注：见正文引文处原注标记（编号可能按页或分节重排）" in comment
    assert "对应脚注编号" not in comment
    assert "66" not in comment
    assert "审核原因：逐字核对一致、校注意见只写入 Word 批注，不改正文" in comment
    assert "exact_text" not in comment

    legacy_label = dict(
        candidate,
        issue_label="已核验原文；管理员确认后仅以 Word 批注写入",
    )
    legacy_comment = public_tasks._candidate_comment_text(legacy_label)
    assert "采信后将作为 Word 批注写入，不改动正文" in legacy_comment
    assert "管理员确认后" not in legacy_comment

    literal = dict(candidate, existing_note_kind="manual_endnote", existing_note_marker="〔43〕")
    literal_comment = public_tasks._candidate_comment_text(literal)
    assert "对应手工尾注标记：〔43〕" in literal_comment

    cross_reference = dict(candidate, existing_note_kind="reference_field")
    cross_comment = public_tasks._candidate_comment_text(cross_reference)
    assert "对应交叉引用：见正文域标记（更新域后可能变化，仅供人工复核）" in cross_comment
    assert "66" not in cross_comment


def test_noteref_cross_reference_is_readonly_and_not_comment_eligible() -> None:
    paragraph = public_tasks.etree.fromstring(
        (
            f'<w:p xmlns:w="{public_tasks.W_NS}"><w:r><w:t>引文正文</w:t></w:r>'
            '<w:r><w:instrText> NOTEREF _Ref987654 \\h </w:instrText>'
            '<w:t>3</w:t></w:r></w:p>'
        ).encode("utf-8"),
        parser=public_tasks._safe_xml_parser(),
    )
    references = public_tasks._paragraph_reference_fields(paragraph, {})
    assert references == [{
        "kind": "reference_field", "id": 987654, "field_target": "_Ref987654",
        "offset": 4, "text": "", "readonly": True,
        "reference_mode": "note_cross_reference",
    }]
    candidate = _accepted_exact_candidate(kind="audit")
    candidate["existing_note_kind"] = "reference_field"
    assert public_tasks._candidate_is_comment_eligible({"mode": "audit"}, candidate) is False


def test_pdf_position_recovery_is_current_version_pdf_only_and_lossless(
    isolated_test_store: Path,
) -> None:
    legacy_error = public_tasks.LEGACY_PDF_POSITION_ERROR_PREFIX + "“社会生活在本质上是实践的”。"

    def failed_pdf_job(
        name: str, *, corpus: str = "corpus", template: str = "template",
        error: str = legacy_error, pdf_status: str = "position_failed",
        keep_word: bool = True, expired: bool = False,
    ) -> tuple[dict, Path]:
        job = test_tasks.create_job(
            7, f"{name}.docx", _docx(isolated_test_store / f"{name}-source.docx"),
            mode="audit", recognition_depth="direct_only", scope_tokens=["book:文集"],
            corpus_sha256=corpus, template_version=template,
        )
        current = test_tasks.get_job(str(job["id"]), 7)
        word_path = Path(str(current["input_path"])).parent / f"{name}-final.docx"
        word_path.write_bytes(Path(str(current["input_path"])).read_bytes())
        test_tasks.replace_candidates(str(job["id"]), [_accepted_exact_candidate(kind="audit")])
        candidate = test_tasks.all_candidates(str(job["id"]))[0]
        test_tasks.save_decisions(
            str(job["id"]), 7,
            [{
                "id": candidate["id"], "decision": "accepted",
                "selected_option": candidate["selected_option"],
                "proposed_citation": candidate["proposed_citation"],
            }],
        )
        test_tasks.update_job(
            str(job["id"]), status="complete", word_export_status="ready",
            pdf_export_status=pdf_status, output_docx_path=str(word_path),
            output_pdf_path=str(word_path.with_suffix(".pdf")), error=error,
            progress_done=1, progress_total=1, pdf_position_failure_count=1,
        )
        if not keep_word:
            word_path.unlink()
        if expired:
            with test_tasks.core._connect() as conn:
                conn.execute(
                    "UPDATE citation_assistant_jobs SET expires_at=? WHERE id=?",
                    ("2000-01-01T00:00:00+00:00", str(job["id"])),
                )
        return job, word_path

    recoverable, word_path = failed_pdf_job("recoverable")
    word_hash = hashlib.sha256(word_path.read_bytes()).hexdigest()
    body_validation, body_word_path = failed_pdf_job(
        "body-validation", error=public_tasks.LEGACY_PDF_BODY_ERROR_PREFIX,
        pdf_status="failed",
    )
    body_word_hash = hashlib.sha256(body_word_path.read_bytes()).hexdigest()
    wrong_template, _ = failed_pdf_job("wrong-template", template="old-template")
    wrong_corpus, _ = failed_pdf_job("wrong-corpus", corpus="old-corpus")
    other_error, _ = failed_pdf_job("other-error", error="Word 已生成；PDF 未生成：转换失败。")
    expired, _ = failed_pdf_job("expired", expired=True)
    missing_word, _ = failed_pdf_job("missing-word", keep_word=False)

    assert test_tasks.recover_pdf_position_failures("corpus", "template") == 2
    current = test_tasks.get_job(str(recoverable["id"]), 7)
    assert current["status"] == "exporting"
    assert current["word_export_status"] == "ready"
    assert current["pdf_export_status"] == "converting"
    assert current["output_docx_path"] == str(word_path)
    assert current["output_pdf_path"] == ""
    assert current["error"] == ""
    assert current["progress_done"] == 0 and current["progress_total"] == 1
    assert hashlib.sha256(word_path.read_bytes()).hexdigest() == word_hash
    assert test_tasks.all_candidates(str(recoverable["id"]))[0]["decision"] == "accepted"
    body_current = test_tasks.get_job(str(body_validation["id"]), 7)
    assert body_current["status"] == "exporting"
    assert body_current["pdf_export_status"] == "converting"
    assert hashlib.sha256(body_word_path.read_bytes()).hexdigest() == body_word_hash
    for unchanged in (wrong_template, wrong_corpus, other_error, expired, missing_word):
        row = test_tasks.get_job(str(unchanged["id"]), 7)
        assert row["status"] == "complete"
        assert row["pdf_export_status"] == "position_failed"


@pytest.mark.parametrize("note_kind", ["footnote", "endnote"])
def test_insertion_word_uses_blue_reference_marker_and_note_text(
    isolated_test_store: Path, note_kind: str,
) -> None:
    job = test_tasks.create_job(
        7, "插注样本.docx", _docx(isolated_test_store / f"insert-{note_kind}.docx"), mode="generate",
        note_kind=note_kind,
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )
    current = test_tasks.get_job(str(job["id"]), 7)
    candidate = _accepted_exact_candidate(kind="generate")
    candidate["writeback_mode"] = note_kind
    output = public_tasks.export_docx(current, [candidate])
    assert output.name.endswith("_插注版.docx")
    with zipfile.ZipFile(output) as package:
        body = public_tasks._parse_xml(package.read("word/document.xml"))
        notes = public_tasks._parse_xml(package.read(f"word/{note_kind}s.xml"))
        body_colors = body.xpath(
            f".//w:{note_kind}Reference/parent::w:r/w:rPr/w:color/@w:val",
            namespaces=public_tasks.NS,
        )
        note_colors = notes.xpath(
            f"./w:{note_kind}[number(@w:id) > 0]//w:color/@w:val",
            namespaces=public_tasks.NS,
        )
        assert body_colors == [public_tasks.INSERTION_BLUE]
        assert note_colors and set(note_colors) == {public_tasks.INSERTION_BLUE}


def test_proofreading_word_preserves_existing_notes_and_adds_body_comment(
    isolated_test_store: Path,
) -> None:
    insertion_job = test_tasks.create_job(
        7, "原有脚注.docx", _docx(isolated_test_store / "source.docx"), mode="generate",
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )
    inserted = public_tasks.export_docx(
        test_tasks.get_job(str(insertion_job["id"]), 7),
        [_accepted_exact_candidate(kind="generate")],
    )
    audit_job = test_tasks.create_job(
        8, "待校注.docx", inserted.read_bytes(), mode="audit",
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )
    audit_current = test_tasks.get_job(str(audit_job["id"]), 8)
    output = public_tasks.export_docx(audit_current, [_accepted_exact_candidate(kind="audit")])
    assert output.name.endswith("_校注批注版.docx")
    with zipfile.ZipFile(audit_current["input_path"]) as original, zipfile.ZipFile(output) as exported:
        assert exported.read("word/footnotes.xml") == original.read("word/footnotes.xml")
        assert "word/comments.xml" in exported.namelist()
        body = public_tasks._parse_xml(exported.read("word/document.xml"))
        assert len(body.xpath(".//w:commentRangeStart", namespaces=public_tasks.NS)) == 1


def test_combined_word_contains_blue_insertion_and_proofreading_comment(
    isolated_test_store: Path,
) -> None:
    job = test_tasks.create_job(
        7, "组合样本.docx", _docx(isolated_test_store / "combined.docx"), mode="both",
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )
    current = test_tasks.get_job(str(job["id"]), 7)
    output = public_tasks.export_docx(current, [
        _accepted_exact_candidate(kind="generate"),
        _accepted_exact_candidate(kind="audit"),
    ])
    assert output.name.endswith("_插注校注版.docx")
    with zipfile.ZipFile(output) as package:
        body = public_tasks._parse_xml(package.read("word/document.xml"))
        assert body.xpath(
            ".//w:footnoteReference/parent::w:r/w:rPr/w:color[@w:val=$blue]",
            namespaces=public_tasks.NS, blue=public_tasks.INSERTION_BLUE,
        )
        assert len(body.xpath(".//w:commentRangeStart", namespaces=public_tasks.NS)) == 1
        assert "word/comments.xml" in package.namelist()


def test_generate_mode_skips_pdf_and_reports_not_requested(
    isolated_test_store: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = test_tasks.create_job(
        7, "插注样本.docx", _docx(isolated_test_store / "word-only-source.docx"), mode="generate",
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )
    test_tasks.update_job(str(job["id"]), status="review_ready")
    word_path = isolated_test_store / "inserted.docx"
    word_path.write_bytes(b"docx")
    calls: list[str] = []
    monkeypatch.setattr(test_tasks.core, "export_docx", lambda *_args, **_kwargs: word_path)
    monkeypatch.setattr(test_tasks.core, "export_pdf", lambda *_args, **_kwargs: calls.append("pdf"))
    test_tasks.run_export(str(job["id"]))
    current = test_tasks.get_job(str(job["id"]), 7)
    assert calls == []
    assert current and current["pdf_export_status"] == "not_requested"
    assert current["output_docx_path"] == str(word_path)


def test_pdf_uses_final_word_and_only_proofreading_manifest(
    isolated_test_store: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = test_tasks.create_job(
        7, "组合样本.docx", _docx(isolated_test_store / "pdf-source.docx"), mode="both",
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )
    current = test_tasks.get_job(str(job["id"]), 7)
    final_word = isolated_test_store / "final-combined.docx"
    final_word.write_bytes(Path(current["input_path"]).read_bytes())
    seen: dict[str, object] = {}

    def fake_convert(source: Path, output: Path) -> None:
        seen["source"] = source
        output.write_bytes(b"%PDF-test")

    def fake_margin(_path: Path, candidates: list[dict]) -> dict:
        seen["candidates"] = candidates
        return {"located": len(candidates), "failed": 0}

    monkeypatch.setattr(public_tasks, "_convert_annotated_docx_to_pdf", fake_convert)
    monkeypatch.setattr(public_tasks, "_render_review_margin_annotations", fake_margin)
    monkeypatch.setattr(public_tasks, "_validate_annotated_pdf", lambda *_args, **_kwargs: {})
    pdf_path = public_tasks.export_pdf(current, [
        _accepted_exact_candidate(kind="generate"),
        _accepted_exact_candidate(kind="audit"),
    ], source_docx=final_word)
    assert pdf_path.name == "组合样本_插注校注批注版.pdf"
    assert seen["source"] == final_word
    assert [item["kind"] for item in seen["candidates"]] == ["audit"]
    audit_job = dict(current)
    audit_job["mode"] = "audit"
    audit_path = public_tasks.export_pdf(
        audit_job, [_accepted_exact_candidate(kind="audit")], source_docx=final_word,
    )
    assert audit_path.name == "组合样本_校注批注版.pdf"


def test_pdf_conversion_copy_removes_only_agent_comments_and_keeps_source_unchanged(
    isolated_test_store: Path,
) -> None:
    source = isolated_test_store / "comments-and-notes.docx"
    document = Document()
    document.add_heading("正文", level=1)
    user_paragraph = document.add_paragraph()
    user_run = user_paragraph.add_run("这是用户自己留下的批注文字。")
    document.add_comment(user_run, text="用户原有批注", author="论文作者", initials="作者")
    document.add_paragraph("马克思指出：“社会生活在本质上是实践的。”")
    document.add_paragraph("马克思指出：“社会生活在本质上是实践的。”")
    document.save(source)
    job = test_tasks.create_job(
        7, source.name, source.read_bytes(), mode="both", note_kind="footnote",
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )
    current = test_tasks.get_job(str(job["id"]), 7)
    final_word = public_tasks.export_docx(current, [
        _accepted_exact_candidate(kind="audit", paragraph_index=2),
        _accepted_exact_candidate(kind="generate", paragraph_index=3),
    ])
    source_hash = hashlib.sha256(final_word.read_bytes()).hexdigest()
    with zipfile.ZipFile(final_word) as package:
        original_notes = package.read("word/footnotes.xml")
        original_comments = public_tasks._parse_xml(package.read("word/comments.xml"))
        assert len(original_comments.xpath("./w:comment", namespaces=public_tasks.NS)) == 2

    conversion_copy = public_tasks._build_pdf_conversion_copy_without_agent_comments(final_word)
    try:
        assert conversion_copy != final_word
        assert hashlib.sha256(final_word.read_bytes()).hexdigest() == source_hash
        with zipfile.ZipFile(conversion_copy) as package:
            comments = public_tasks._parse_xml(package.read("word/comments.xml"))
            body = public_tasks._parse_xml(package.read("word/document.xml"))
            remaining = comments.xpath("./w:comment", namespaces=public_tasks.NS)
            assert len(remaining) == 1
            assert remaining[0].get(f"{{{public_tasks.W_NS}}}author") == "论文作者"
            assert remaining[0].get(f"{{{public_tasks.W_NS}}}initials") == "作者"
            user_id = remaining[0].get(f"{{{public_tasks.W_NS}}}id")
            for marker in ("commentRangeStart", "commentRangeEnd", "commentReference"):
                ids = body.xpath(f".//w:{marker}/@w:id", namespaces=public_tasks.NS)
                assert ids == [user_id]
            assert package.read("word/footnotes.xml") == original_notes
    finally:
        conversion_copy.unlink(missing_ok=True)
    assert not list(final_word.parent.glob(".*.pdf-source.docx"))


@pytest.mark.parametrize("outcome", ["success", "failure", "timeout"])
def test_agent_comment_conversion_copy_is_cleaned_for_every_outcome(
    isolated_test_store: Path, monkeypatch: pytest.MonkeyPatch, outcome: str,
) -> None:
    job = test_tasks.create_job(
        7, "清理样本.docx", _docx(isolated_test_store / "cleanup-source.docx"), mode="audit",
        recognition_depth="direct_only", scope_tokens=["book:文集"],
        corpus_sha256="corpus", template_version="template",
    )
    current = test_tasks.get_job(str(job["id"]), 7)
    candidate = _accepted_exact_candidate(kind="audit")
    final_word = public_tasks.export_docx(current, [candidate])
    source_hash = hashlib.sha256(final_word.read_bytes()).hexdigest()
    seen_sources: list[Path] = []

    def fake_convert(conversion_source: Path, output: Path) -> None:
        seen_sources.append(Path(conversion_source))
        assert Path(conversion_source) != final_word
        with zipfile.ZipFile(conversion_source) as package:
            comments = public_tasks._parse_xml(package.read("word/comments.xml"))
            assert not comments.xpath("./w:comment", namespaces=public_tasks.NS)
        if outcome == "failure":
            raise public_tasks.CitationAssistantError("转换失败")
        if outcome == "timeout":
            raise public_tasks.CitationAssistantError("批注式 PDF 转换超时")
        output.write_bytes(b"%PDF-test")

    monkeypatch.setattr(public_tasks, "_convert_annotated_docx_to_pdf", fake_convert)
    monkeypatch.setattr(
        public_tasks, "_render_review_margin_annotations",
        lambda _path, items: {"located": len(items), "failed": 0},
    )
    monkeypatch.setattr(public_tasks, "_validate_annotated_pdf", lambda *_args, **_kwargs: {})
    if outcome == "success":
        public_tasks.export_pdf(
            current, [candidate], source_docx=final_word,
            suppress_agent_comments_for_conversion=True,
        )
    else:
        with pytest.raises(public_tasks.CitationAssistantError):
            public_tasks.export_pdf(
                current, [candidate], source_docx=final_word,
                suppress_agent_comments_for_conversion=True,
            )
    assert len(seen_sources) == 1
    assert not seen_sources[0].exists()
    assert not list(final_word.parent.glob(".*.pdf-source.docx"))
    assert not list(final_word.parent.glob(".*.pdf-source.docx.part"))
    assert hashlib.sha256(final_word.read_bytes()).hexdigest() == source_hash


def test_pdf_locator_rejects_missing_or_ambiguous_comment_anchor(
    isolated_test_store: Path,
) -> None:
    path = isolated_test_store / "locator.pdf"
    document = public_tasks.fitz.open()
    for _index in range(2):
        page = document.new_page()
        page.insert_text((72, 72), "unique quotation anchor and repeated quotation anchor")
    document.save(path)
    document.close()
    with public_tasks.fitz.open(path) as pdf:
        geometries = [public_tasks._pdf_normalized_geometry(page) for page in pdf]
        page_texts = [geometry[0] for geometry in geometries]
        with pytest.raises(public_tasks.CitationAssistantError, match="存在多个候选位置"):
            public_tasks._locate_pdf_candidate(
                pdf, page_texts, {"paper_text": "repeated quotation anchor"},
                page_geometries=geometries,
            )
        with pytest.raises(public_tasks.CitationAssistantError, match="文本层未找到可验证位置"):
            public_tasks._locate_pdf_candidate(
                pdf, page_texts, {"paper_text": "quotation that is absent"},
                page_geometries=geometries,
            )


def test_pdf_locator_joins_normalized_character_boxes_across_lines(tmp_path: Path) -> None:
    path = tmp_path / "wrapped-locator.pdf"
    document = public_tasks.fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "quotation anchor starts")
    page.insert_text((72, 92), "and continues across lines")
    document.save(path)
    document.close()
    candidate = {"paper_text": "quotation anchor starts and continues across lines"}
    with public_tasks.fitz.open(path) as pdf:
        geometries = [public_tasks._pdf_normalized_geometry(page) for page in pdf]
        page_texts = [geometry[0] for geometry in geometries]
        page_index, rectangles = public_tasks._locate_pdf_candidate(
            pdf, page_texts, candidate, page_geometries=geometries,
        )
    assert page_index == 0
    assert len(rectangles) >= 2
    assert len({round(rectangle.y0) for rectangle in rectangles}) == 2


def test_pdf_locator_sorts_out_of_storage_order_fragments_by_page_coordinates(
    tmp_path: Path,
) -> None:
    path = tmp_path / "visual-order-locator.pdf"
    document = public_tasks.fitz.open()
    page = document.new_page()
    # Insert the middle fragment first to reproduce LibreOffice's nonvisual
    # PDF object order while keeping the rendered line left-to-right.
    page.insert_text((230, 72), "middle anchor")
    page.insert_text((72, 72), "visual beginning")
    page.insert_text((340, 72), "closing phrase")
    document.save(path)
    document.close()

    candidate = {"paper_text": "visual beginning middle anchor closing phrase"}
    with public_tasks.fitz.open(path) as pdf:
        assert public_tasks.normalize(pdf[0].get_text("text")) != public_tasks.normalize(candidate["paper_text"])
        geometry = public_tasks._pdf_normalized_geometry(pdf[0])
        assert geometry[0] == public_tasks.normalize(candidate["paper_text"])
        page_index, rectangles = public_tasks._locate_pdf_candidate(
            pdf, [geometry[0]], candidate, page_geometries=[geometry],
        )
    assert page_index == 0
    assert rectangles
    assert rectangles[0].x0 < rectangles[-1].x0


def test_pdf_validation_uses_visual_order_for_body_retention(tmp_path: Path) -> None:
    source = tmp_path / "visual-order-source.docx"
    body_text = "visual beginning middle anchor closing phrase retained body text"
    document = Document()
    document.add_paragraph(body_text)
    document.save(source)

    pdf_path = tmp_path / "visual-order-validated.pdf"
    pdf = public_tasks.fitz.open()
    page = pdf.new_page()
    page.insert_text((230, 72), "middle anchor")
    page.insert_text((72, 72), "visual beginning")
    page.insert_text((340, 72), "closing phrase retained body text")
    page.insert_text((72, 102), "proofreading note")
    pdf.save(pdf_path)
    pdf.close()

    with public_tasks.fitz.open(pdf_path) as check:
        assert public_tasks.normalize(check[0].get_text("text")) != public_tasks.normalize(body_text)
    result = public_tasks._validate_annotated_pdf(
        pdf_path,
        {"input_path": str(source)},
        [{"decision": "accepted", "issue_label": "proofreading note"}],
    )
    assert result == {"pages": 1, "rendered_pages": 1}
